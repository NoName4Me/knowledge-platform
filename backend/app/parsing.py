from __future__ import annotations

import html
import io
import re
import shutil
import subprocess
import tempfile
import zipfile
from dataclasses import dataclass
from pathlib import Path

from docx import Document as DocxDocument
from docx.oxml.ns import qn
from docx.table import Table, _Row
from docx.text.paragraph import Paragraph
from pypdf import PdfReader

SUPPORTED_EXTENSIONS = (".pdf", ".doc", ".docx", ".txt", ".md", ".html", ".htm")
SUPPORTED_FORMAT_LABEL = "PDF、Word（DOC/DOCX）、TXT、Markdown 和 HTML"

CHUNK_CHAR_LIMIT = 1200
CHUNK_OVERLAP = 150
MAX_EXTRACTED_CHARS = 2_000_000
MAX_CHUNKS = 5_000
DOCX_MAX_EXPANDED_BYTES = 100 * 1024 * 1024
DOCX_MAX_ENTRIES = 10_000
DOCX_MAX_COMPRESSION_RATIO = 100


class UnsupportedDocument(Exception):
    pass


@dataclass
class ExtractedChunk:
    content: str
    locator_label: str
    locator_start: int | None = None
    locator_end: int | None = None


@dataclass
class ExtractedDocument:
    chunks: list[ExtractedChunk]
    title: str


def extension_of(filename: str) -> str:
    dot = filename.rfind(".")
    return filename[dot:].lower() if dot != -1 else ""


def _split_text(text: str, label_prefix: str) -> list[ExtractedChunk]:
    normalized = re.sub(r"[ \t\r\n]+", " ", text).strip()
    if not normalized:
        return []

    chunks: list[ExtractedChunk] = []
    start = 0
    ordinal = 0
    while start < len(normalized):
        end = min(start + CHUNK_CHAR_LIMIT, len(normalized))
        piece = normalized[start:end].strip()
        if piece:
            ordinal += 1
            chunks.append(
                ExtractedChunk(
                    content=piece,
                    locator_label=f"{label_prefix} {ordinal}",
                    locator_start=start,
                    locator_end=end,
                )
            )
        if end >= len(normalized):
            break
        start = max(end - CHUNK_OVERLAP, start + 1)
    return chunks


def _normalize_plain_text(text: str) -> str:
    # 保留段落边界，避免把两个独立段落拼成同一引用位置。
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    lines = [line.rstrip() for line in text.splitlines()]
    return "\n".join(lines).strip()


def _extract_pdf(data: bytes) -> list[ExtractedChunk]:
    if not data.lstrip().startswith(b"%PDF-"):
        raise UnsupportedDocument("文件声明为 PDF，但内容不是有效 PDF")
    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        raise UnsupportedDocument("PDF 已加密，无法提取文字")
    chunks: list[ExtractedChunk] = []
    for page_number, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        for chunk in _split_text(text, "片段"):
            # 与纯文本、Markdown 保持一致的“位置 · 片段序号”读法，
            # 避免出现“第 1 页 · 第 1 页 1”这类重复前缀。
            chunk.locator_label = f"第 {page_number} 页 · {chunk.locator_label}"
            chunks.append(chunk)
    return chunks


def _extract_legacy_doc(
    data: bytes,
    converter: str = "soffice",
    timeout_seconds: int = 60,
) -> list[ExtractedChunk]:
    """Use an isolated LibreOffice profile to convert legacy Word to DOCX."""
    command = shutil.which(converter) or converter
    if not Path(command).is_file():
        raise UnsupportedDocument(
            "旧版 Word（.doc）需要 LibreOffice 转换器；"
            "请安装 LibreOffice 或配置 KP_LEGACY_DOC_CONVERTER"
        )
    if not data.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
        raise UnsupportedDocument("文件声明为 DOC，但内容不是有效旧版 Word 文档")

    with tempfile.TemporaryDirectory(prefix="kp-doc-") as temp:
        root = Path(temp)
        source = root / "source.doc"
        output = root / "output"
        profile = root / "profile"
        output.mkdir()
        source.write_bytes(data)
        try:
            completed = subprocess.run(
                [
                    command,
                    "--headless",
                    f"-env:UserInstallation={profile.as_uri()}",
                    "--convert-to",
                    "docx",
                    "--outdir",
                    str(output),
                    str(source),
                ],
                capture_output=True,
                check=False,
                text=True,
                timeout=timeout_seconds,
            )
        except subprocess.TimeoutExpired as exc:
            raise UnsupportedDocument("旧版 Word 转换超时") from exc
        except OSError as exc:
            raise UnsupportedDocument("旧版 Word 转换器启动失败") from exc
        converted = output / "source.docx"
        if completed.returncode != 0 or not converted.is_file():
            detail = (completed.stderr or completed.stdout or "").strip()
            suffix = f"：{detail[:180]}" if detail else ""
            raise UnsupportedDocument(f"旧版 Word 转换失败，请确认文件未损坏{suffix}")
        return _extract_docx(converted.read_bytes())


def _extract_docx(data: bytes) -> list[ExtractedChunk]:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            infos = archive.infolist()
            if len(infos) > DOCX_MAX_ENTRIES:
                raise UnsupportedDocument("DOCX 内部条目过多")
            total_size = sum(info.file_size for info in infos)
            if total_size > DOCX_MAX_EXPANDED_BYTES:
                raise UnsupportedDocument("DOCX 解压后体积超限")
            if any(
                info.compress_size > 0
                and info.file_size / info.compress_size > DOCX_MAX_COMPRESSION_RATIO
                for info in infos
            ):
                raise UnsupportedDocument("DOCX 压缩比异常")
    except zipfile.BadZipFile as exc:
        raise UnsupportedDocument("文件声明为 DOCX，但内容不是有效 DOCX") from exc
    document = DocxDocument(io.BytesIO(data))
    chunks: list[ExtractedChunk] = []
    paragraph_number = 0
    table_number = 0
    for block in document.element.body.iterchildren():
        if block.tag == qn("w:p"):
            paragraph_number += 1
            paragraph = Paragraph(block, document)
            for chunk in _split_text(paragraph.text, "片段"):
                chunk.locator_label = f"第 {paragraph_number} 段 · {chunk.locator_label}"
                chunks.append(chunk)
        elif block.tag == qn("w:tbl"):
            table_number += 1
            table = Table(block, document)
            chunks.extend(_extract_docx_table(table, table_number))
    return chunks


def _row_texts(row: _Row) -> list[str]:
    """读取一行单元格文字，合并单元格在 python-docx 中只计一次。"""
    values: list[str] = []
    seen: set[int] = set()
    for cell in row.cells:
        marker = id(cell._tc)
        if marker in seen:
            continue
        seen.add(marker)
        value = re.sub(r"[ \t\r\n]+", " ", cell.text).strip()
        if value:
            values.append(value)
    return values


def _extract_docx_table(table: Table, table_number: int) -> list[ExtractedChunk]:
    """保留普通表格的单元格文字；不尝试还原合并、跨页或嵌套表格结构。"""
    rows = [_row_texts(row) for row in table.rows]
    rows = [row for row in rows if row]
    if not rows:
        return []
    header = rows[0] if len(rows) > 1 else []
    chunks: list[ExtractedChunk] = []
    for row_number, values in enumerate(rows, start=1):
        if header and len(header) == len(values):
            content = "；".join(
                f"{name}：{value}"
                for name, value in zip(header, values, strict=True)
            )
        else:
            content = " | ".join(values)
        for chunk in _split_text(content, "片段"):
            chunk.locator_label = (
                f"第 {table_number} 张表 · 第 {row_number} 行 · {chunk.locator_label}"
            )
            chunks.append(chunk)
    return chunks


def _extract_plain(data: bytes) -> list[ExtractedChunk]:
    try:
        text = data.decode("utf-8")
    except UnicodeDecodeError as exc:
        raise UnsupportedDocument("文本文件不是 UTF-8 编码") from exc
    return _extract_text(text)


def _extract_html(data: bytes) -> list[ExtractedChunk]:
    try:
        text = data.decode("utf-8")
    except UnicodeDecodeError as exc:
        raise UnsupportedDocument("HTML 网页不是 UTF-8 编码") from exc
    # 首版只保留读者可见正文，不执行脚本，也不跟踪页面链接。
    text = re.sub(r"(?is)<(script|style|noscript|svg|template).*?</\1>", " ", text)
    text = re.sub(r"(?i)<br\s*/?>", "\n", text)
    text = re.sub(r"(?i)</(p|div|section|article|li|h[1-6]|tr)>", "\n\n", text)
    text = re.sub(r"(?s)<[^>]+>", " ", text)
    text = html.unescape(text)
    return _extract_text(text)


def _extract_text(text: str) -> list[ExtractedChunk]:
    blocks = [block for block in re.split(r"\n\s*\n", _normalize_plain_text(text)) if block]
    chunks: list[ExtractedChunk] = []
    for block_number, block in enumerate(blocks, start=1):
        for chunk in _split_text(block, "片段"):
            chunk.locator_label = f"第 {block_number} 段 · {chunk.locator_label}"
            chunks.append(chunk)
    return chunks


def extract_document(
    filename: str,
    data: bytes,
    *,
    legacy_doc_converter: str = "soffice",
    legacy_doc_timeout: int = 60,
) -> ExtractedDocument:
    extension = extension_of(filename)
    if extension not in SUPPORTED_EXTENSIONS:
        raise UnsupportedDocument(
            f"暂不支持 {extension or '该'} 格式，当前支持 {SUPPORTED_FORMAT_LABEL}"
        )

    if extension == ".pdf":
        chunks = _extract_pdf(data)
    elif extension == ".docx":
        chunks = _extract_docx(data)
    elif extension == ".doc":
        chunks = _extract_legacy_doc(data, legacy_doc_converter, legacy_doc_timeout)
    elif extension in {".html", ".htm"}:
        chunks = _extract_html(data)
    else:
        chunks = _extract_plain(data)

    if not chunks:
        raise UnsupportedDocument("没有提取到可检索的文字，可能是扫描件或空白文件")
    if sum(len(chunk.content) for chunk in chunks) > MAX_EXTRACTED_CHARS:
        raise UnsupportedDocument("提取文字超过 200 万字符限制")
    if len(chunks) > MAX_CHUNKS:
        raise UnsupportedDocument("资料片段超过 5000 个限制")

    title = filename.rsplit("/", 1)[-1].rsplit(".", 1)[0]
    return ExtractedDocument(chunks=chunks, title=title)
