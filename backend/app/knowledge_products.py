"""V3 personal notebook and report services.

The report is a saved deliverable, not another analysis view. Section edits never
write back to an analysis. Source snapshots let a report remain explainable even
after a source is deleted; current availability is evaluated separately.
"""

from __future__ import annotations

import io
import json
import re
import uuid
from collections import defaultdict
from typing import Any, cast

import httpx
from docx import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import func, select, update
from sqlalchemy.orm import Session

from app.analysis import source_fingerprint
from app.auth import TeamContext
from app.config import get_settings
from app.knowledge_base_access import require_knowledge_base
from app.models import (
    AnalysisCitation,
    AnalysisConclusion,
    AnalysisDocument,
    AnalysisField,
    AnalysisFieldResult,
    AnalysisTask,
    Document,
    DocumentChunk,
    KnowledgeAutomation,
    KnowledgeBase,
    KnowledgeNote,
    Report,
    ReportCitation,
    ReportSection,
    ReportSource,
    SavedRequirement,
)
from app.schemas import (
    AnalysisFieldInput,
    AnalysisSuggestionRequest,
    EvidenceSnapshotInput,
    EvidenceSnapshotOut,
    MissingItemOut,
    NoteCreateRequest,
    NoteOut,
    NoteUpdateRequest,
    ReportFromAnalysisRequest,
    ReportFromSourcesRequest,
    ReportGapOut,
    ReportOut,
    ReportSectionOut,
    ReportSectionSuggestionOut,
    ReportSectionSuggestionRequest,
    ReportSourceOut,
    ReportSummaryOut,
    ReportUpdateRequest,
    SavedRequirementCreate,
)

settings = get_settings()


class KnowledgeProductError(Exception):
    def __init__(self, detail: str, status_code: int = 422) -> None:
        super().__init__(detail)
        self.detail = detail
        self.status_code = status_code


class ProductNotFound(KnowledgeProductError):
    def __init__(self, detail: str = "记录不存在") -> None:
        super().__init__(detail, 404)


class ProductConflict(KnowledgeProductError):
    def __init__(self, detail: str) -> None:
        super().__init__(detail, 409)


def _knowledge_base(session: Session, context: TeamContext, kb_id: uuid.UUID) -> KnowledgeBase:
    try:
        knowledge_base, _permission = require_knowledge_base(session, context, kb_id)
    except HTTPException as exc:
        if exc.status_code == 403:
            raise ProductNotFound("知识库不存在或无权访问") from exc
        raise ProductNotFound("知识库不存在") from exc
    return knowledge_base


def _snapshot(citation: EvidenceSnapshotInput | dict[str, Any]) -> dict[str, Any]:
    if isinstance(citation, EvidenceSnapshotInput):
        payload = citation.model_dump(mode="json")
    else:
        payload = dict(citation)
    return {
        "document_id": payload.get("document_id"),
        "document_title": str(payload.get("document_title") or "资料"),
        "locator_label": str(payload.get("locator_label") or ""),
        "excerpt": str(payload.get("excerpt") or ""),
    }


def _note_out(note: KnowledgeNote) -> NoteOut:
    return NoteOut(
        id=note.id,
        knowledge_base_id=note.knowledge_base_id,
        title=note.title,
        content=note.content,
        source_type=note.source_type,
        source_id=note.source_id,
        citations=[
            EvidenceSnapshotOut.model_validate(item) for item in note.citations
        ],
        created_at=note.created_at,
        updated_at=note.updated_at,
    )


def list_notes(
    session: Session, context: TeamContext, kb_id: uuid.UUID
) -> list[KnowledgeNote]:
    _knowledge_base(session, context, kb_id)
    return list(
        session.scalars(
            select(KnowledgeNote)
            .where(
                KnowledgeNote.team_id == context.team.id,
                KnowledgeNote.knowledge_base_id == kb_id,
                KnowledgeNote.created_by_id == context.user.id,
            )
            .order_by(KnowledgeNote.updated_at.desc(), KnowledgeNote.id.desc())
        )
    )


def create_note(
    session: Session, context: TeamContext, kb_id: uuid.UUID, payload: NoteCreateRequest
) -> NoteOut:
    _knowledge_base(session, context, kb_id)
    note = KnowledgeNote(
        team_id=context.team.id,
        knowledge_base_id=kb_id,
        created_by_id=context.user.id,
        title=payload.title.strip(),
        content=payload.content.strip(),
        source_type=payload.source_type,
        source_id=payload.source_id,
        citations=[_snapshot(item) for item in payload.citations],
    )
    session.add(note)
    session.flush()
    return _note_out(note)


def _note(
    session: Session, context: TeamContext, kb_id: uuid.UUID, note_id: uuid.UUID
) -> KnowledgeNote:
    _knowledge_base(session, context, kb_id)
    note = session.scalar(
        select(KnowledgeNote)
        .where(
            KnowledgeNote.id == note_id,
            KnowledgeNote.team_id == context.team.id,
            KnowledgeNote.knowledge_base_id == kb_id,
            KnowledgeNote.created_by_id == context.user.id,
        )
        .with_for_update()
    )
    if note is None:
        raise ProductNotFound("笔记不存在")
    return note


def get_note(
    session: Session, context: TeamContext, kb_id: uuid.UUID, note_id: uuid.UUID
) -> NoteOut:
    return _note_out(_note(session, context, kb_id, note_id))


def update_note(
    session: Session,
    context: TeamContext,
    kb_id: uuid.UUID,
    note_id: uuid.UUID,
    payload: NoteUpdateRequest,
) -> NoteOut:
    note = _note(session, context, kb_id, note_id)
    note.title = payload.title.strip()
    note.content = payload.content.strip()
    note.citations = [_snapshot(item) for item in payload.citations]
    session.flush()
    return _note_out(note)


def delete_note(
    session: Session, context: TeamContext, kb_id: uuid.UUID, note_id: uuid.UUID
) -> None:
    note = _note(session, context, kb_id, note_id)
    session.delete(note)
    session.flush()


def _analysis(session: Session, context: TeamContext, analysis_id: uuid.UUID) -> AnalysisTask:
    task = session.scalar(
        select(AnalysisTask).where(
            AnalysisTask.id == analysis_id,
            AnalysisTask.team_id == context.team.id,
            AnalysisTask.created_by_id == context.user.id,
        )
    )
    if task is None:
        raise ProductNotFound("分析任务不存在")
    return task


def _available_document(session: Session, context: TeamContext, doc_id: uuid.UUID) -> Document:
    document = session.scalar(
        select(Document).where(
            Document.id == doc_id,
            Document.team_id == context.team.id,
            Document.deleted_at.is_(None),
            Document.status == "ready",
        )
    )
    if document is None:
        raise KnowledgeProductError("所选资料不可用")
    return document


def _analysis_payload(
    session: Session, context: TeamContext, task: AnalysisTask
) -> tuple[dict[str, Any], list[dict[str, Any]], bool, str | None]:
    """Return a displayable analysis snapshot and its available-state warning."""
    executions = list(
        session.scalars(
            select(AnalysisDocument)
            .where(AnalysisDocument.task_id == task.id)
            .order_by(AnalysisDocument.ordinal)
        )
    )
    if task.status not in {"completed", "partial"}:
        raise ProductConflict("只有已完成的本人分析才能生成报告")
    from app import analysis as analysis_service

    invalid_reason = analysis_service.source_invalid_reason(session, executions)
    if invalid_reason is not None:
        raise ProductConflict(f"分析来源已失效：{invalid_reason}")
    fields = list(
        session.scalars(
            select(AnalysisField)
            .where(AnalysisField.task_id == task.id)
            .order_by(AnalysisField.ordinal)
        )
    )
    fields_by_id = {field.id: field for field in fields}
    titles: dict[uuid.UUID, str] = {}
    for execution in executions:
        source_document = session.get(Document, execution.document_id)
        titles[execution.document_id] = (
            source_document.title if source_document is not None else "已删除资料"
        )
    citations = list(
        session.scalars(
            select(AnalysisCitation)
            .where(AnalysisCitation.task_id == task.id)
            .order_by(AnalysisCitation.ordinal)
        )
    )
    by_target: dict[uuid.UUID, list[dict[str, Any]]] = defaultdict(list)
    for citation in citations:
        target = citation.field_result_id or citation.conclusion_id
        if target is None:
            continue
        by_target[target].append(
            {
                "document_id": citation.document_id,
                "document_title": titles.get(citation.document_id, "资料"),
                "locator_label": citation.locator_label,
                "excerpt": citation.excerpt,
                "table_number": citation.table_number,
                "table_row": citation.table_row,
                "table_headers": citation.table_headers,
                "table_cells": citation.table_cells,
            }
        )
    results = list(
        session.scalars(
            select(AnalysisFieldResult)
            .where(AnalysisFieldResult.task_id == task.id)
            .order_by(AnalysisFieldResult.field_id, AnalysisFieldResult.id)
        )
    )
    execution_by_id: dict[uuid.UUID, AnalysisDocument] = {
        item.id: item for item in executions
    }
    result_rows = []
    for result in results:
        result_execution = execution_by_id.get(result.document_execution_id)
        field = fields_by_id.get(result.field_id)
        if (
            result_execution is None
            or field is None
            or result_execution.status != "succeeded"
        ):
            continue
        value = (
            result.raw_value
            or result.value_text
            or (
                str(result.value_number) + (f" {result.unit}" if result.unit else "")
                if result.value_number is not None
                else None
            )
            or ("、".join(result.value_list) if result.value_list else None)
            or ("未提及" if result.status == "not_mentioned" else "待核实")
        )
        result_rows.append(
            {
                "document_title": titles.get(result_execution.document_id, "资料"),
                "field_name": field.name,
                "value": value,
                "status": result.status,
                "citations": by_target.get(result.id, []),
            }
        )
    conclusions = list(
        session.scalars(
            select(AnalysisConclusion)
            .where(AnalysisConclusion.task_id == task.id)
            .order_by(AnalysisConclusion.ordinal)
        )
    )
    return (
        {
            "task": task,
            "fields": fields,
            "results": result_rows,
            "conclusions": conclusions,
            "citations": by_target,
            "documents": [
                {"id": item.document_id, "title": titles.get(item.document_id, "资料")}
                for item in executions
            ],
        },
        [
            {
                "source_type": "document",
                "source_id": item.document_id,
                "title_snapshot": titles.get(item.document_id, "资料"),
                "fingerprint": source_fingerprint(
                    [
                        (row.id, row.locator_label, row.content)
                        for row in session.scalars(
                            select(DocumentChunk)
                            .where(DocumentChunk.document_id == item.document_id)
                            .order_by(DocumentChunk.ordinal)
                        )
                    ]
                ),
            }
            for item in executions
        ],
        task.status == "partial",
        "部分来源分析未完成，报告仅包含已完成资料。" if task.status == "partial" else None,
    )


def _add_section(
    session: Session,
    *,
    report: Report,
    ordinal: int,
    heading: str,
    body: str,
    citations: list[dict[str, Any]] | None = None,
) -> ReportSection:
    section = ReportSection(
        team_id=report.team_id,
        knowledge_base_id=report.knowledge_base_id,
        report_id=report.id,
        owner_id=report.created_by_id,
        ordinal=ordinal,
        heading=heading,
        body=body,
    )
    session.add(section)
    session.flush()
    seen: set[tuple[str, str, str]] = set()
    for item in citations or []:
        if not item.get("excerpt"):
            continue
        key = (
            str(item.get("document_id") or ""),
            str(item.get("locator_label") or ""),
            str(item.get("excerpt") or ""),
        )
        if key in seen:
            continue
        seen.add(key)
        session.add(
            ReportCitation(
                team_id=report.team_id,
                knowledge_base_id=report.knowledge_base_id,
                report_id=report.id,
                owner_id=report.created_by_id,
                section_id=section.id,
                ordinal=len(seen) - 1,
                document_id=item.get("document_id"),
                document_title=str(item.get("document_title") or "资料")[:300],
                locator_label=str(item.get("locator_label") or "")[:200],
                excerpt=str(item.get("excerpt") or "")[:4000],
            )
        )
    return section


def _analysis_sections(
    payload: dict[str, Any], focus: str
) -> list[tuple[str, str, list[dict[str, Any]]]]:
    groups: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in payload["results"]:
        groups[row["field_name"]].append(row)
    sections: list[tuple[str, str, list[dict[str, Any]]]] = []
    overview: list[str] = []
    citations: list[dict[str, Any]] = []
    for document in payload["documents"]:
        overview.append(f"- {document['title']}")
    for field in payload["fields"]:
        rows = groups.get(field.name, [])
        if not rows:
            continue
        lines = []
        section_citations: list[dict[str, Any]] = []
        for row in rows:
            lines.append(f"- {row['document_title']}：{row['value']}")
            section_citations.extend(row["citations"])
        sections.append(
            (field.name, "\n".join(lines), section_citations)
        )
    conclusion_lines = []
    conclusion_citations = []
    for conclusion in payload["conclusions"]:
        label = {"common": "共同点", "difference": "差异", "needs_review": "待核实"}[
            conclusion.kind
        ]
        conclusion_lines.append(f"- {label}：{conclusion.text}")
        conclusion_citations.extend(payload["citations"].get(conclusion.id, []))
    if conclusion_lines:
        sections.append(("主要发现", "\n".join(conclusion_lines), conclusion_citations))
    if not sections:
        sections.append(
            (
                "分析结果",
                focus or "本次分析没有形成可写入报告的字段结果。",
                citations,
            )
        )
    sections.insert(
        0,
        (
            "资料范围",
            "\n".join(overview) if overview else "未记录资料范围。",
            citations,
        ),
    )
    return sections


def _report_out(session: Session, report: Report) -> ReportOut:
    sections = list(
        session.scalars(
            select(ReportSection)
            .where(ReportSection.report_id == report.id)
            .order_by(ReportSection.ordinal)
        )
    )
    citations = list(
        session.scalars(
            select(ReportCitation)
            .where(ReportCitation.report_id == report.id)
            .order_by(ReportCitation.section_id, ReportCitation.ordinal)
        )
    )
    by_section: dict[uuid.UUID, list[ReportCitation]] = defaultdict(list)
    for citation in citations:
        by_section[citation.section_id].append(citation)
    sources = list(
        session.scalars(
            select(ReportSource)
            .where(ReportSource.report_id == report.id)
            .order_by(ReportSource.created_at, ReportSource.id)
        )
    )
    available: dict[str, bool] = {}
    for source in sources:
        if source.source_type == "document":
            document = session.get(Document, source.source_id)
            available[str(source.source_id)] = bool(
                document is not None
                and document.deleted_at is None
                and document.status == "ready"
            )
        elif source.source_type == "note":
            available[str(source.source_id)] = (
                session.scalar(
                    select(KnowledgeNote.id).where(
                        KnowledgeNote.id == source.source_id,
                        KnowledgeNote.team_id == report.team_id,
                        KnowledgeNote.created_by_id == report.created_by_id,
                    )
                )
                is not None
            )
        elif source.source_type == "analysis":
            task = session.get(AnalysisTask, source.source_id)
            # 删除分析后来源快照仍用于解释报告出处；ID 已解除关联时视为已删除。
            available[str(source.source_id)] = task is not None and task.status in {
                "completed",
                "partial",
            }
        else:
            available[str(source.source_id)] = False
    return ReportOut(
        id=report.id,
        knowledge_base_id=report.knowledge_base_id,
        source_analysis_id=report.source_analysis_id,
        title=report.title,
        purpose=report.purpose,
        focus=report.focus,
        length=report.length,
        status=report.status,
        needs_review=report.needs_review,
        review_note=report.review_note,
        source_review_required=report.source_review_required,
        source_review_note=report.source_review_note,
        gaps=[
            MissingItemOut.model_validate(item)
            for item in (report.gaps or [])
            if isinstance(item, dict) and item.get("text")
        ],
        sections=[
            ReportSectionOut(
                id=section.id,
                ordinal=section.ordinal,
                heading=section.heading,
                body=section.body,
                citations=[
                    {
                        "id": citation.id,
                        "ordinal": citation.ordinal,
                        "document_id": citation.document_id,
                        "document_title": citation.document_title,
                        "locator_label": citation.locator_label,
                        "excerpt": citation.excerpt,
                    }
                    for citation in by_section.get(section.id, [])
                ],
            )
            for section in sections
        ],
        sources=[
            ReportSourceOut(
                id=source.id,
                source_type=source.source_type,
                source_id=source.source_id,
                title_snapshot=source.title_snapshot,
                available=available.get(str(source.source_id), True),
            )
            for source in sources
        ],
        created_at=report.created_at,
        updated_at=report.updated_at,
    )


def _source_rows(
    session: Session, context: TeamContext, document_ids: list[uuid.UUID], note_ids: list[uuid.UUID]
) -> tuple[list[Document], list[KnowledgeNote]]:
    if len(document_ids) != len(set(document_ids)) or len(note_ids) != len(set(note_ids)):
        raise KnowledgeProductError("来源清单中存在重复项目")
    if not document_ids and not note_ids:
        raise KnowledgeProductError("请至少选择一份资料或一条笔记")
    documents = [
        _available_document(session, context, doc_id) for doc_id in document_ids
    ]
    notes = [
        session.scalar(
            select(KnowledgeNote).where(
                KnowledgeNote.id == note_id,
                KnowledgeNote.team_id == context.team.id,
                KnowledgeNote.created_by_id == context.user.id,
            )
        )
        for note_id in note_ids
    ]
    if any(note is None for note in notes):
        raise ProductNotFound("所选笔记不存在")
    return documents, [note for note in notes if note is not None]


def _source_report_evidence(
    session: Session,
    documents: list[Document],
    notes: list[KnowledgeNote],
) -> tuple[list[dict[str, Any]], list[str], bool]:
    """Build a bounded evidence list; model output may only cite these entries."""
    evidence: list[dict[str, Any]] = []
    note_context: list[str] = []
    seen: set[tuple[str, str, str]] = set()
    budget = max(settings.analysis_batch_chars * 6, 24_000)
    used = 0
    truncated = False
    for document in documents:
        chunks = list(
            session.scalars(
                select(DocumentChunk)
                .where(DocumentChunk.document_id == document.id)
                .order_by(DocumentChunk.ordinal)
            )
        )
        for chunk in chunks:
            if used + len(chunk.content) > budget:
                truncated = True
                break
            evidence.append(
                {
                    "document_id": document.id,
                    "document_title": document.title,
                    "locator_label": chunk.locator_label,
                    "excerpt": chunk.content[:4000],
                }
            )
            used += len(chunk.content)
        if truncated:
            break
    for note in notes:
        note_context.append(
            f"笔记《{note.title}》（这是用户理解或摘录，不是原始证据）：\n{note.content[:4000]}"
        )
        for item in note.citations:
            snapshot = _snapshot(item)
            key = (
                str(snapshot["document_id"] or ""),
                snapshot["locator_label"],
                snapshot["excerpt"],
            )
            if key in seen:
                continue
            seen.add(key)
            evidence.append(snapshot)
    return evidence, note_context, truncated


def _generate_source_report_with_model(
    *,
    title: str,
    purpose: str,
    focus: str,
    length: str,
    documents: list[Document],
    notes: list[KnowledgeNote],
    evidence: list[dict[str, Any]],
    note_context: list[str],
    truncated: bool,
) -> tuple[list[tuple[str, str, list[dict[str, Any]]]], list[dict[str, str]]] | None:
    if not _has_report_model():
        return None
    numbered = "\n\n".join(
        f"[{index}] {item['document_title']}（{item['locator_label']}）\n"
        f"{item['excerpt']}"
        for index, item in enumerate(evidence, start=1)
    )
    notes_text = "\n\n".join(note_context) or "未选择笔记。"
    length_guide = {
        "short": "约 800–1200 字，2–4 个章节",
        "medium": "约 1500–2500 字，3–6 个章节",
        "long": "约 3000–5000 字，5–8 个章节",
    }[length]
    prompt = (
        "根据已选资料起草一份可编辑报告。只能使用资料证据中的事实，"
        "不得引入外部知识、补全未写明的数字或伪造原文。笔记只帮助理解用户已有观点，"
        "不能把笔记中的推断当作原始事实。"
        "返回 JSON：{\"sections\":[{\"heading\":\"...\",\"body\":\"...\","
        "\"citation_numbers\":[1,2]}]}。每个事实性章节使用 citation_numbers "
        "指向下方证据编号；没有依据的章节可留空数组，但不得编造依据。"
        "正文使用清晰段落或短列表，不写 Markdown 标题，不输出 JSON 之外的内容。\n\n"
        f"报告名称：{title}\n用途：{purpose or '未指定'}\n重点：{focus or '未指定'}\n"
        f"篇幅：{length_guide}\n\n"
        f"资料证据：\n{numbered}\n\n用户笔记：\n{notes_text}"
    )
    try:
        response = httpx.post(
            f"{settings.analysis_model_base_url.rstrip('/')}/chat/completions",
            headers={"Authorization": f"Bearer {settings.analysis_model_api_key}"},
            json={
                "model": settings.analysis_model,
                "messages": [
                    {
                        "role": "system",
                        "content": "你是严谨的中文报告撰稿人，只依据给定证据写作。",
                    },
                    {"role": "user", "content": prompt},
                ],
                "temperature": 0.2,
                "response_format": {"type": "json_object"},
            },
            timeout=min(settings.analysis_model_timeout_seconds, 120),
        )
        response.raise_for_status()
        content = response.json()["choices"][0]["message"]["content"]
        parsed = json.loads(content)
    except (httpx.HTTPError, ValueError, KeyError, IndexError, TypeError) as exc:
        raise KnowledgeProductError(
            "报告生成模型暂时不可用，已保留所选来源，请稍后重试", 502
        ) from exc
    raw_sections = parsed.get("sections") if isinstance(parsed, dict) else None
    if not isinstance(raw_sections, list) or not raw_sections:
        raise KnowledgeProductError("报告生成模型没有返回可用章节，请稍后重试", 502)
    sections: list[tuple[str, str, list[dict[str, Any]]]] = []
    for raw in raw_sections[:12]:
        if not isinstance(raw, dict):
            continue
        heading = str(raw.get("heading") or "").strip()[:200]
        body = str(raw.get("body") or "").strip()[:100_000]
        if not heading or not body:
            continue
        citations: list[dict[str, Any]] = []
        raw_numbers = raw.get("citation_numbers")
        if isinstance(raw_numbers, list):
            for value in raw_numbers[:50]:
                if isinstance(value, int) and 1 <= value <= len(evidence):
                    citations.append(evidence[value - 1])
        sections.append((heading, body, citations))
    if not sections:
        raise KnowledgeProductError("报告生成模型返回的章节未通过校验，请稍后重试", 502)
    gaps: list[dict[str, str]] = [
        {
            "kind": "review",
            "text": "报告由所选资料与笔记生成；提交前请逐章核对事实、条件和引用。",
        }
    ]
    if truncated:
        gaps.append(
            {
                "kind": "source",
                "text": "本次生成只纳入可处理范围内的资料正文；未覆盖部分不能视为已核对。",
            }
        )
    if notes:
        gaps.append(
            {
                "kind": "review",
                "text": "笔记中的理解和归纳已用于组织报告，但不能替代原始资料证据。",
            }
        )
    return sections, gaps


def create_report_from_analysis(
    session: Session,
    context: TeamContext,
    kb_id: uuid.UUID,
    analysis_id: uuid.UUID,
    payload: ReportFromAnalysisRequest,
) -> ReportOut:
    _knowledge_base(session, context, kb_id)
    task = _analysis(session, context, analysis_id)
    if task.knowledge_base_id != kb_id:
        raise ProductNotFound("分析任务不存在")
    data, sources, needs_review, warning = _analysis_payload(session, context, task)
    title = (payload.title or "").strip() or f"{task.title} · 报告"
    report = Report(
        team_id=context.team.id,
        knowledge_base_id=kb_id,
        created_by_id=context.user.id,
        source_analysis_id=task.id,
        title=title[:200],
        purpose=payload.purpose.strip(),
        focus=payload.focus.strip(),
        length=payload.length,
        status="draft",
        needs_review=needs_review,
        review_note=warning,
        gaps=[
            {
                "kind": "missing",
                "text": (
                    f"“{field.name}”没有任何成功提取的资料结果，"
                    "请确认是否需要补充材料。"
                ),
            }
            for field in data["fields"]
            if not any(row["field_name"] == field.name for row in data["results"])
        ]
        + [
            {"kind": "review", "text": "部分字段存在未提及或待核实结果，请在提交前逐项核对。"}
            for _ in [0]
            if any(row["status"] in {"not_mentioned", "needs_review"} for row in data["results"])
        ],
    )
    session.add(report)
    session.flush()
    for source in sources:
        session.add(
            ReportSource(
                team_id=context.team.id,
                knowledge_base_id=kb_id,
                report_id=report.id,
                owner_id=context.user.id,
                **source,
            )
        )
    session.add(
        ReportSource(
            team_id=context.team.id,
            knowledge_base_id=kb_id,
            report_id=report.id,
            owner_id=context.user.id,
            source_type="analysis",
            source_id=task.id,
            title_snapshot=task.title,
            fingerprint=task.request_fingerprint,
        )
    )
    for ordinal, (heading, body, citations) in enumerate(_analysis_sections(data, payload.focus)):
        _add_section(
            session,
            report=report,
            ordinal=ordinal,
            heading=heading,
            body=body,
            citations=citations,
        )
    session.flush()
    return _report_out(session, report)


def create_report_from_sources(
    session: Session,
    context: TeamContext,
    kb_id: uuid.UUID,
    payload: ReportFromSourcesRequest,
) -> ReportOut:
    _knowledge_base(session, context, kb_id)
    documents, notes = _source_rows(
        session, context, payload.document_ids, payload.note_ids
    )
    for document in documents:
        if document.knowledge_base_id != kb_id:
            raise KnowledgeProductError("所选资料不属于当前知识库")
    for note in notes:
        if note.knowledge_base_id != kb_id:
            raise KnowledgeProductError("所选笔记不属于当前知识库")
    report = Report(
        team_id=context.team.id,
        knowledge_base_id=kb_id,
        created_by_id=context.user.id,
        title=payload.title.strip(),
        purpose=payload.purpose.strip(),
        focus=payload.focus.strip(),
        length=payload.length,
        gaps=[],
    )
    session.add(report)
    session.flush()
    overview = []
    citations: list[dict[str, Any]] = []
    for document in documents:
        chunks = list(
            session.scalars(
                select(DocumentChunk)
                .where(DocumentChunk.document_id == document.id)
                .order_by(DocumentChunk.ordinal)
                .limit(3)
            )
        )
        overview.append(f"- {document.title}")
        citations.extend(
            {
                "document_id": document.id,
                "document_title": document.title,
                "locator_label": chunk.locator_label,
                "excerpt": chunk.content[:500],
            }
            for chunk in chunks
        )
        session.add(
            ReportSource(
                team_id=context.team.id,
                knowledge_base_id=kb_id,
                report_id=report.id,
                owner_id=context.user.id,
                source_type="document",
                source_id=document.id,
                title_snapshot=document.title,
                fingerprint=source_fingerprint(
                    [
                        (row.id, row.locator_label, row.content)
                        for row in session.scalars(
                            select(DocumentChunk)
                            .where(DocumentChunk.document_id == document.id)
                            .order_by(DocumentChunk.ordinal)
                        )
                    ]
                ),
            )
        )
    note_bodies = []
    for note in notes:
        note_bodies.append(note.content)
        overview.append(f"- 笔记：{note.title}")
        citations.extend(_snapshot(item) for item in note.citations)
        session.add(
            ReportSource(
                team_id=context.team.id,
                knowledge_base_id=kb_id,
                report_id=report.id,
                owner_id=context.user.id,
                source_type="note",
                source_id=note.id,
                title_snapshot=note.title,
            )
        )
    evidence, note_context, truncated = _source_report_evidence(
        session, documents, notes
    )
    generated = _generate_source_report_with_model(
        title=report.title,
        purpose=report.purpose,
        focus=report.focus,
        length=report.length,
        documents=documents,
        notes=notes,
        evidence=evidence,
        note_context=note_context,
        truncated=truncated,
    )
    if generated is not None:
        generated_sections, gaps = generated
        report.gaps = cast(list[dict[str, object]], gaps)
        for ordinal, (heading, body, section_citations) in enumerate(generated_sections):
            _add_section(
                session,
                report=report,
                ordinal=ordinal,
                heading=heading,
                body=body,
                citations=section_citations,
            )
        session.flush()
        return _report_out(session, report)
    report.gaps = [
        {
            "kind": "source",
            "text": "当前正文来自资料与笔记的初始摘录，请核对是否还需补充关键结论。",
        }
    ]
    _add_section(
        session,
        report=report,
        ordinal=0,
        heading="资料与笔记",
        body="\n".join(overview) or "没有可用来源。",
        citations=citations,
    )
    _add_section(
        session,
        report=report,
        ordinal=1,
        heading="主要内容",
        body="\n\n".join(note_bodies)
        or "所选资料已加入报告来源。请结合右侧原文核对后补写主要发现。",
        citations=citations if not note_bodies else [],
    )
    session.flush()
    return _report_out(session, report)


def list_reports(
    session: Session, context: TeamContext, kb_id: uuid.UUID
) -> list[ReportSummaryOut]:
    _knowledge_base(session, context, kb_id)
    rows = session.execute(
        select(Report, func.count(ReportSection.id))
        .outerjoin(ReportSection, ReportSection.report_id == Report.id)
        .where(
            Report.team_id == context.team.id,
            Report.knowledge_base_id == kb_id,
            Report.created_by_id == context.user.id,
        )
        .group_by(Report.id)
        .order_by(Report.updated_at.desc(), Report.id.desc())
    ).all()
    return [
        ReportSummaryOut(
            id=report.id,
            title=report.title,
            status=report.status,
            needs_review=report.needs_review,
            source_review_required=report.source_review_required,
            source_analysis_id=report.source_analysis_id,
            section_count=count,
            created_at=report.created_at,
            updated_at=report.updated_at,
        )
        for report, count in rows
    ]


def _report(
    session: Session, context: TeamContext, kb_id: uuid.UUID, report_id: uuid.UUID
) -> Report:
    _knowledge_base(session, context, kb_id)
    report = session.scalar(
        select(Report)
        .where(
            Report.id == report_id,
            Report.team_id == context.team.id,
            Report.knowledge_base_id == kb_id,
            Report.created_by_id == context.user.id,
        )
        .with_for_update()
    )
    if report is None:
        raise ProductNotFound("报告不存在")
    return report


def get_report(
    session: Session, context: TeamContext, kb_id: uuid.UUID, report_id: uuid.UUID
) -> ReportOut:
    return _report_out(session, _report(session, context, kb_id, report_id))


def update_report(
    session: Session,
    context: TeamContext,
    kb_id: uuid.UUID,
    report_id: uuid.UUID,
    payload: ReportUpdateRequest,
) -> ReportOut:
    report = _report(session, context, kb_id, report_id)
    report.title = payload.title.strip()
    report.purpose = payload.purpose.strip()
    report.focus = payload.focus.strip()
    report.length = payload.length
    report.status = payload.status
    if payload.source_review_acknowledged and report.source_review_required:
        report.source_review_required = False
        report.source_review_note = None
    report.gaps = [item.model_dump(mode="json") for item in payload.gaps]
    existing = {
        section.id: section
        for section in session.scalars(
            select(ReportSection).where(ReportSection.report_id == report.id)
        )
    }
    # Move retained sections out of the user-visible ordinal range first. Swapping
    # two sections otherwise violates the per-report unique ordinal constraint.
    for existing_section in existing.values():
        existing_section.ordinal += 1000
    session.flush()
    seen: set[uuid.UUID] = set()
    for ordinal, section_payload in enumerate(payload.sections):
        saved_section = (
            existing.get(section_payload.id) if section_payload.id else None
        )
        if saved_section is None:
            saved_section = ReportSection(
                team_id=report.team_id,
                knowledge_base_id=report.knowledge_base_id,
                report_id=report.id,
                owner_id=report.created_by_id,
                heading=section_payload.heading.strip(),
                body=section_payload.body.strip(),
                ordinal=ordinal,
            )
            session.add(saved_section)
            session.flush()
            continue
        seen.add(saved_section.id)
        saved_section.heading = section_payload.heading.strip()
        saved_section.body = section_payload.body.strip()
        saved_section.ordinal = ordinal
    for section_id, existing_section in existing.items():
        if section_id not in seen:
            session.delete(existing_section)
    session.flush()
    return _report_out(session, report)


def delete_report(
    session: Session, context: TeamContext, kb_id: uuid.UUID, report_id: uuid.UUID
) -> None:
    report = _report(session, context, kb_id, report_id)
    # 复合外键不能只把 report_id 置空而不影响团队/知识库/创建者，因此先解除关联。
    session.execute(
        update(KnowledgeAutomation)
        .where(
            KnowledgeAutomation.team_id == context.team.id,
            KnowledgeAutomation.knowledge_base_id == kb_id,
            KnowledgeAutomation.created_by_id == context.user.id,
            KnowledgeAutomation.report_id == report.id,
        )
        .values(report_id=None)
    )
    session.delete(report)
    session.flush()


def list_requirements(
    session: Session, context: TeamContext, kb_id: uuid.UUID, kind: str | None = None
) -> list[SavedRequirement]:
    _knowledge_base(session, context, kb_id)
    stmt = select(SavedRequirement).where(
        SavedRequirement.team_id == context.team.id,
        SavedRequirement.knowledge_base_id == kb_id,
        SavedRequirement.created_by_id == context.user.id,
    )
    if kind:
        stmt = stmt.where(SavedRequirement.kind == kind)
    return list(
        session.scalars(
            stmt.order_by(
                SavedRequirement.updated_at.desc(), SavedRequirement.id.desc()
            )
        )
    )


def create_requirement(
    session: Session, context: TeamContext, kb_id: uuid.UUID, payload: SavedRequirementCreate
) -> SavedRequirement:
    _knowledge_base(session, context, kb_id)
    existing = session.scalar(
        select(SavedRequirement).where(
            SavedRequirement.team_id == context.team.id,
            SavedRequirement.created_by_id == context.user.id,
            SavedRequirement.kind == payload.kind,
            SavedRequirement.name == payload.name.strip(),
        )
    )
    if existing is not None:
        existing.knowledge_base_id = kb_id
        existing.payload = payload.payload
        session.flush()
        return existing
    requirement = SavedRequirement(
        team_id=context.team.id,
        knowledge_base_id=kb_id,
        created_by_id=context.user.id,
        kind=payload.kind,
        name=payload.name.strip(),
        payload=payload.payload,
    )
    session.add(requirement)
    session.flush()
    return requirement


def delete_requirement(
    session: Session, context: TeamContext, kb_id: uuid.UUID, requirement_id: uuid.UUID
) -> None:
    _knowledge_base(session, context, kb_id)
    row = session.scalar(
        select(SavedRequirement).where(
            SavedRequirement.id == requirement_id,
            SavedRequirement.team_id == context.team.id,
            SavedRequirement.knowledge_base_id == kb_id,
            SavedRequirement.created_by_id == context.user.id,
        )
    )
    if row is None:
        raise ProductNotFound("常用要求不存在")
    session.delete(row)
    session.flush()


def suggest_analysis_fields(
    session: Session,
    context: TeamContext,
    kb_id: uuid.UUID,
    payload: AnalysisSuggestionRequest,
) -> tuple[list[AnalysisFieldInput], str]:
    _knowledge_base(session, context, kb_id)
    documents = [_available_document(session, context, item) for item in payload.document_ids]
    if any(document.knowledge_base_id != kb_id for document in documents):
        raise KnowledgeProductError("所选资料不属于当前知识库")
    if settings.analysis_model_base_url and settings.analysis_model_api_key:
        result = _suggest_with_model(payload.requirement, documents, session)
        if result is not None:
            return result
    # Deterministic fallback used only when no model is configured or the request cannot
    # be parsed. It keeps the UI useful without pretending semantic certainty.
    base = payload.requirement.strip()
    chunks: list[DocumentChunk] = []
    for document in documents[:2]:
        chunks.extend(
            session.scalars(
                select(DocumentChunk)
                .where(DocumentChunk.document_id == document.id)
                .order_by(DocumentChunk.ordinal)
                .limit(1)
            )
        )
    context_text = " ".join(chunk.content[:500] for chunk in chunks)
    fields = [
        AnalysisFieldInput(
            name="用途或目标",
            description=f"围绕“{base}”识别材料说明的目的、用途或要解决的问题",
            value_type="text",
        ),
        AnalysisFieldInput(
            name="关键条件",
            description="列出适用范围、前提条件、限制或例外",
            value_type="text_list",
        ),
    ]
    if re.search(r"时间|日期|何时|期限", base):
        fields.append(
            AnalysisFieldInput(
                name="关键时间",
                description="提取生效、截止或其他与要求相关的时间",
                value_type="text_list",
            )
        )
    elif re.search(r"金额|价格|费用|成本", base):
        fields.append(
            AnalysisFieldInput(
                name="相关数值",
                description="提取与要求直接相关的金额、比例或数量",
                value_type="text_list",
            )
        )
    rationale = (
        "已按你描述的用途给出可调整字段。"
        + (" 当前未配置真实分析模型，字段由本地规则给出。" if not context_text else "")
    )
    return fields[:6], rationale


def _suggest_with_model(
    requirement: str, documents: list[Document], session: Session
) -> tuple[list[AnalysisFieldInput], str] | None:
    excerpts = []
    for document in documents[:3]:
        chunk = session.scalar(
            select(DocumentChunk)
            .where(DocumentChunk.document_id == document.id)
            .order_by(DocumentChunk.ordinal)
        )
        if chunk is not None:
            excerpts.append(f"{document.title}：{chunk.content[:800]}")
    prompt = (
        "根据用户目的建议 2-6 个信息提取字段。只返回 JSON："
        '{"fields":[{"name":"...","description":"...","value_type":"text|number|date|text_list"}],'
        '"rationale":"..."}。字段名称简短，说明明确；不要引用外部事实。\n\n'
        f"用户目的：{requirement}\n资料样例：\n" + "\n".join(excerpts)
    )
    try:
        response = httpx.post(
            f"{settings.analysis_model_base_url.rstrip('/')}/chat/completions",
            headers={"Authorization": f"Bearer {settings.analysis_model_api_key}"},
            json={
                "model": settings.analysis_model,
                "messages": [
                    {"role": "system", "content": "你是文档分析字段设计助手。"},
                    {"role": "user", "content": prompt},
                ],
                "temperature": 0,
                "response_format": {"type": "json_object"},
            },
            timeout=min(settings.analysis_model_timeout_seconds, 45),
        )
        response.raise_for_status()
        body = response.json()
        content = body["choices"][0]["message"]["content"]
        parsed = json.loads(content)
        fields = [
            AnalysisFieldInput.model_validate(item)
            for item in parsed.get("fields", [])
        ]
        if fields:
            rationale = str(parsed.get("rationale") or "以下字段已根据你的目的整理。")
            return fields[:6], rationale[:1000]
    except (httpx.HTTPError, ValueError, KeyError, TypeError):
        return None
    return None


def document_guide(
    session: Session, context: TeamContext, kb_id: uuid.UUID, document_id: uuid.UUID
) -> dict[str, Any]:
    _knowledge_base(session, context, kb_id)
    document = _available_document(session, context, document_id)
    if document.knowledge_base_id != kb_id:
        raise ProductNotFound("资料不存在")
    chunks = list(
        session.scalars(
            select(DocumentChunk)
            .where(DocumentChunk.document_id == document.id)
            .order_by(DocumentChunk.ordinal)
        )
    )
    if not chunks:
        raise ProductNotFound("资料没有可读正文")
    joined = "\n".join(chunk.content for chunk in chunks)
    sentences = [item.strip() for item in re.split(r"[。！？\n]+", joined) if item.strip()]
    summary = " ".join(sentences[:4])[:700] if sentences else joined[:700]
    labels = list(dict.fromkeys(chunk.locator_label for chunk in chunks))
    terms = re.findall(r"[\u4e00-\u9fff]{2,8}", joined)
    stop = {"可以", "以及", "进行", "如果", "对于", "这个", "其中", "相关", "应当", "不得"}
    counts: dict[str, int] = {}
    for term in terms:
        if term in stop or len(term) < 2:
            continue
        counts[term] = counts.get(term, 0) + 1
    concepts = [
        term
        for term, count in sorted(counts.items(), key=lambda item: (-item[1], -len(item[0])))
        if count > 1
    ][:8]
    questions = [
        "这份资料解决什么问题？",
        "有哪些适用范围、条件和限制？",
        "资料中的关键结论或要求是什么？",
        "哪些内容仍需要其他材料补充？",
    ]
    return {
        "document_id": document.id,
        "document_title": document.title,
        "summary": summary or "正文暂无足够内容形成摘要。",
        "section_labels": labels[:20],
        "key_concepts": concepts,
        "suggested_questions": questions,
        "coverage_note": (
            f"导读覆盖已解析的 {len(chunks)} 个正文片段；"
            "扫描件、图表及复杂版式可能未包含。"
        ),
    }


def search_documents(
    session: Session,
    context: TeamContext,
    kb_id: uuid.UUID,
    query: str,
    document_id: uuid.UUID | None,
) -> list[dict[str, Any]]:
    _knowledge_base(session, context, kb_id)
    normalized = query.strip()
    if not normalized:
        raise KnowledgeProductError("搜索内容不能为空")
    stmt = (
        select(DocumentChunk, Document)
        .join(Document, Document.id == DocumentChunk.document_id)
        .where(
            DocumentChunk.team_id == context.team.id,
            Document.knowledge_base_id == kb_id,
            Document.deleted_at.is_(None),
            Document.status == "ready",
            DocumentChunk.content.ilike(f"%{normalized}%"),
        )
    )
    if document_id is not None:
        stmt = stmt.where(Document.id == document_id)
    rows = list(session.execute(stmt.order_by(Document.title, DocumentChunk.ordinal).limit(40)))
    results = []
    for chunk, document in rows:
        position = chunk.content.lower().find(normalized.lower())
        start = max(0, position - 180)
        end = min(len(chunk.content), position + len(normalized) + 300)
        results.append(
            {
                "document_id": document.id,
                "document_title": document.title,
                "locator_label": chunk.locator_label,
                "excerpt": chunk.content[position : position + 500],
                "before": chunk.content[start:position],
                "after": chunk.content[position + len(normalized) : end],
            }
        )
    return results


def report_gaps(
    session: Session, context: TeamContext, kb_id: uuid.UUID, report_id: uuid.UUID
) -> ReportGapOut:
    report = _report_out(session, _report(session, context, kb_id, report_id))
    items = list(report.gaps)
    for source in report.sources:
        if not source.available:
            items.append(
                MissingItemOut(
                    kind="source",
                    text=f"保存时使用的{source.title_snapshot}已不可用；正文仍保留，但不能证明当前内容。",
                )
            )
    if not items:
        items.append(
            MissingItemOut(
                kind="review",
                text="当前没有记录到明确缺口；提交前仍需核对关键事实和引用。",
            )
        )
    return ReportGapOut(
        items=items,
        coverage_note="清单来自当前报告、字段结果和来源状态；它不证明已完成全局完备性检查。",
    )


def suggest_report_section(
    session: Session,
    context: TeamContext,
    kb_id: uuid.UUID,
    report_id: uuid.UUID,
    section_id: uuid.UUID,
    payload: ReportSectionSuggestionRequest,
) -> ReportSectionSuggestionOut:
    report = _report(session, context, kb_id, report_id)
    section = session.scalar(
        select(ReportSection).where(
            ReportSection.id == section_id,
            ReportSection.report_id == report.id,
        )
    )
    if section is None:
        raise ProductNotFound("报告章节不存在")
    citations = list(
        session.scalars(
            select(ReportCitation)
            .where(ReportCitation.section_id == section.id)
            .order_by(ReportCitation.ordinal)
        )
    )
    provider = "rule"
    suggestion = section.body
    rationale = "已按你的修改要求保留当前结构；请采用后再次核对事实与来源。"
    if citations and _has_report_model():
        try:
            result = _suggest_section_with_model(section, citations, payload.instruction)
            if result is not None:
                suggestion, rationale = result
                provider = "http"
        except (httpx.HTTPError, ValueError, KeyError, TypeError):
            suggestion = section.body
    elif citations:
        # Useful deterministic refinement: retain cited excerpts as a review scaffold
        # rather than pretending to rewrite prose without a semantic model.
        suggestion = (
            section.body.rstrip()
            + "\n\n核对提示：\n"
            + "\n".join(f"- {item.excerpt[:300]}" for item in citations[:3])
        )
        rationale = "当前未配置真实生成模型；候选仅加入已保存依据作为核对提示。"
    return ReportSectionSuggestionOut(
        original=section.body,
        suggestion=suggestion,
        rationale=rationale,
        provider=provider,
    )


def _has_report_model() -> bool:
    return bool(settings.analysis_model_base_url and settings.analysis_model_api_key)


def _suggest_section_with_model(
    section: ReportSection, citations: list[ReportCitation], instruction: str
) -> tuple[str, str] | None:
    evidence = "\n".join(
        f"[{item.ordinal + 1}] {item.document_title}（{item.locator_label}）：{item.excerpt}"
        for item in citations
    )
    response = httpx.post(
        f"{settings.analysis_model_base_url.rstrip('/')}/chat/completions",
        headers={"Authorization": f"Bearer {settings.analysis_model_api_key}"},
        json={
            "model": settings.analysis_model,
            "messages": [
                {
                    "role": "system",
                    "content": (
                        "你是报告编辑。只能依据给定原文修改章节，不得补写外部事实。"
                        "返回 JSON：{\"suggestion\":\"...\",\"rationale\":\"...\"}。"
                    ),
                },
                {
                    "role": "user",
                    "content": (
                        f"修改要求：{instruction}\n章节：{section.heading}\n当前正文：{section.body}\n"
                        f"原文依据：\n{evidence}"
                    ),
                },
            ],
            "temperature": 0.2,
            "response_format": {"type": "json_object"},
        },
        timeout=min(settings.analysis_model_timeout_seconds, 60),
    )
    response.raise_for_status()
    content = response.json()["choices"][0]["message"]["content"]
    parsed = json.loads(content)
    suggestion = str(parsed.get("suggestion") or "").strip()
    if not suggestion:
        return None
    return suggestion[:100_000], str(parsed.get("rationale") or "已按修改要求生成候选。")[:1000]


def report_docx(
    session: Session, context: TeamContext, kb_id: uuid.UUID, report_id: uuid.UUID
) -> bytes:
    report = _report_out(session, _report(session, context, kb_id, report_id))
    document = DocxDocument()
    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10.5)
    title = document.add_heading(report.title, level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if report.purpose:
        document.add_paragraph(f"用途：{report.purpose}")
    if report.focus:
        document.add_paragraph(f"重点：{report.focus}")
    for section in report.sections:
        document.add_heading(section.heading, level=1)
        for line in section.body.splitlines():
            if line.strip():
                document.add_paragraph(line)
        if section.citations:
            document.add_heading("参考来源", level=2)
            for citation in section.citations:
                document.add_paragraph(
                    f"[{citation.ordinal + 1}] {citation.document_title} · "
                    f"{citation.locator_label}\n{citation.excerpt}",
                    style="List Number",
                )
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def report_xlsx(
    session: Session, context: TeamContext, kb_id: uuid.UUID, report_id: uuid.UUID
) -> bytes:
    report = _report_out(session, _report(session, context, kb_id, report_id))
    workbook = Workbook()
    sheet = workbook.active
    if sheet is None:
        raise KnowledgeProductError("无法创建工作簿", 500)
    sheet.title = "报告"
    headers = ["章节", "正文", "来源", "位置", "摘录"]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="233B4A")
    for section in report.sections:
        if not section.citations:
            sheet.append([section.heading, section.body, "", "", ""])
        for citation in section.citations:
            sheet.append(
                [
                    section.heading,
                    section.body,
                    citation.document_title,
                    citation.locator_label,
                    citation.excerpt,
                ]
            )
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.column_dimensions["A"].width = 24
    sheet.column_dimensions["B"].width = 70
    sheet.column_dimensions["C"].width = 28
    sheet.column_dimensions["D"].width = 20
    sheet.column_dimensions["E"].width = 60
    sheet.freeze_panes = "A2"
    gap_sheet = workbook.create_sheet("待核实与缺口")
    gap_sheet.append(["类型", "事项"])
    for item in report.gaps:
        gap_sheet.append([item.kind, item.text])
    gap_sheet.column_dimensions["A"].width = 18
    gap_sheet.column_dimensions["B"].width = 100
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def report_pdf(
    session: Session, context: TeamContext, kb_id: uuid.UUID, report_id: uuid.UUID
) -> bytes:
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    report = _report_out(session, _report(session, context, kb_id, report_id))
    stream = io.BytesIO()
    doc = SimpleDocTemplate(
        stream,
        pagesize=A4,
        leftMargin=20 * mm,
        rightMargin=20 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
    )
    styles = getSampleStyleSheet()
    font_name = "STSong-Light"
    try:
        pdfmetrics.getFont(font_name)
    except KeyError:
        pdfmetrics.registerFont(UnicodeCIDFont(font_name))
    title = ParagraphStyle(
        "KPTitle", parent=styles["Title"], fontName=font_name, fontSize=20, alignment=TA_CENTER
    )
    heading = ParagraphStyle(
        "KPHeading", parent=styles["Heading2"], fontName=font_name, fontSize=14, spaceAfter=6
    )
    body_style = ParagraphStyle(
        "KPBody", parent=styles["BodyText"], fontName=font_name, fontSize=10, leading=17
    )
    story = [Paragraph(report.title, title), Spacer(1, 8)]
    if report.purpose:
        story.append(Paragraph(f"用途：{report.purpose}", body_style))
    if report.focus:
        story.append(Paragraph(f"重点：{report.focus}", body_style))
    for section in report.sections:
        story.append(Paragraph(section.heading, heading))
        for line in section.body.splitlines():
            if line.strip():
                story.append(Paragraph(line.replace("&", "&amp;").replace("<", "&lt;"), body_style))
        if section.citations:
            story.append(Paragraph("参考来源", heading))
            for citation in section.citations:
                story.append(
                    Paragraph(
                        f"[{citation.ordinal + 1}] {citation.document_title} · "
                        f"{citation.locator_label}<br/>{citation.excerpt}",
                        body_style,
                    )
                )
    if report.gaps:
        story.append(Paragraph("待核实与材料缺口", heading))
        for item in report.gaps:
            story.append(Paragraph(f"• {item.text}", body_style))
    doc.build(story)
    return stream.getvalue()


def report_csv(
    session: Session, context: TeamContext, kb_id: uuid.UUID, report_id: uuid.UUID
) -> str:
    report = _report_out(session, _report(session, context, kb_id, report_id))
    import csv

    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["章节", "正文", "来源", "位置", "摘录"])
    for section in report.sections:
        if not section.citations:
            writer.writerow([section.heading, section.body, "", "", ""])
        for citation in section.citations:
            writer.writerow(
                [
                    section.heading,
                    section.body,
                    citation.document_title,
                    citation.locator_label,
                    citation.excerpt,
                ]
            )
    return output.getvalue()
