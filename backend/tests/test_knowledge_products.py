from __future__ import annotations

import io
from contextlib import contextmanager

import httpx
from docx import Document as DocxDocument
from openpyxl import load_workbook
from pypdf import PdfReader
from sqlalchemy import select

from tests.conftest import add_member, make_user
from tests.test_analysis_api import (
    _create,
    _run_until_finished,
    _team_with_kb,
    _upload_ready,
)


def _analysis_report(client, team_id, kb_id, document_id):
    analysis = _create(client, team_id, kb_id, document_id).json()
    return analysis


def test_note_crud_is_private_and_preserves_evidence_snapshot(client, session):
    user, team, kb_id = _team_with_kb(client, session, "note-owner@example.com")
    base = f"/api/teams/{team.id}/knowledge-bases/{kb_id}/notes"
    created = client.post(
        base,
        json={
            "title": "制度理解",
            "content": "我确认：住宿上限是每晚 500 元。",
            "source_type": "manual",
            "citations": [
                {
                    "document_id": None,
                    "document_title": "差旅制度",
                    "locator_label": "第 2 段",
                    "excerpt": "住宿费用上限为每晚 500 元。",
                }
            ],
        },
    )
    assert created.status_code == 201, created.text
    note = created.json()
    assert note["citations"][0]["excerpt"] == "住宿费用上限为每晚 500 元。"
    updated = client.put(
        f"{base}/{note['id']}",
        json={
            "title": "制度理解（已核对）",
            "content": "人工补充：申请前需确认城市等级。",
            "citations": note["citations"],
        },
    )
    assert updated.status_code == 200
    assert updated.json()["title"] == "制度理解（已核对）"

    other = make_user(session, "note-other@example.com")
    add_member(session, team, other)
    client.post("/api/auth/login", json={"email": other.email, "password": "password123"})
    assert client.get(base).json()["items"] == []
    assert client.get(f"{base}/{note['id']}").status_code == 404
    assert (
        client.put(
            f"{base}/{note['id']}",
            json={"title": "越权", "content": "不能修改", "citations": []},
        ).status_code
        == 404
    )


def test_report_from_analysis_edit_reopen_and_export(client, session):
    _, team, kb_id = _team_with_kb(client, session)
    document = _upload_ready(
        client,
        session,
        team.id,
        kb_id,
        "差旅制度：交通费用凭有效票据报销。\n\n住宿费用上限为每晚 500 元。",
    )
    analysis = _analysis_report(client, team.id, kb_id, document["id"])
    assert _run_until_finished(session, analysis["id"]) == "completed"
    base = f"/api/teams/{team.id}/knowledge-bases/{kb_id}/reports"
    created = client.post(
        f"{base}/from-analysis/{analysis['id']}",
        json={"purpose": "提交团队审阅", "focus": "交通与住宿", "length": "medium"},
    )
    assert created.status_code == 201, created.text
    report = created.json()
    assert report["source_analysis_id"] == analysis["id"]
    assert len(report["sections"]) >= 2
    assert any(section["citations"] for section in report["sections"])
    assert any(
        source["source_type"] == "document" for source in report["sources"]
    )

    sections = report["sections"]
    sections[1]["heading"] = "我的核对结论"
    sections[1]["body"] = "人工确认：住宿费用按原文上限执行。"
    sections.reverse()
    saved = client.put(
        f"{base}/{report['id']}",
        json={
            "title": "差旅制度审阅稿",
            "purpose": report["purpose"],
            "focus": report["focus"],
            "length": "medium",
            "status": "reviewed",
            "sections": [
                {
                    "id": section["id"],
                    "heading": section["heading"],
                    "body": section["body"],
                }
                for section in sections
            ],
        },
    )
    assert saved.status_code == 200, saved.text
    saved_body = saved.json()
    assert [item["heading"] for item in saved_body["sections"]] == [
        item["heading"] for item in sections
    ]
    reopened = client.get(f"{base}/{report['id']}")
    assert reopened.status_code == 200
    assert reopened.json()["title"] == "差旅制度审阅稿"
    assert reopened.json()["status"] == "reviewed"
    edited_section = next(
        item
        for item in reopened.json()["sections"]
        if item["heading"] == "我的核对结论"
    )
    assert edited_section["body"] == "人工确认：住宿费用按原文上限执行。"

    exported = client.get(f"{base}/{report['id']}/export.docx")
    assert exported.status_code == 200
    assert exported.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument"
    )
    docx = DocxDocument(io.BytesIO(exported.content))
    text = "\n".join(p.text for p in docx.paragraphs)
    assert "差旅制度审阅稿" in text
    assert "人工确认：住宿费用按原文上限执行。" in text
    assert "参考来源" in text

    csv_export = client.get(f"{base}/{report['id']}/export.csv")
    assert csv_export.status_code == 200
    assert "人工确认：住宿费用按原文上限执行。" in csv_export.content.decode("utf-8-sig")

    xlsx_export = client.get(f"{base}/{report['id']}/export.xlsx")
    assert xlsx_export.status_code == 200
    workbook = load_workbook(io.BytesIO(xlsx_export.content))
    assert workbook.sheetnames == ["报告", "待核实与缺口"]
    assert "人工确认：住宿费用按原文上限执行。" in "\n".join(
        str(cell.value or "") for row in workbook["报告"].iter_rows() for cell in row
    )

    pdf_export = client.get(f"{base}/{report['id']}/export.pdf")
    assert pdf_export.status_code == 200
    assert pdf_export.headers["content-type"] == "application/pdf"
    assert len(PdfReader(io.BytesIO(pdf_export.content)).pages) >= 1
    assert client.delete(f"{base}/{report['id']}").status_code == 204
    assert client.get(f"{base}/{report['id']}").status_code == 404


def test_report_keeps_saved_content_when_source_is_deleted(client, session):
    _, team, kb_id = _team_with_kb(client, session)
    document = _upload_ready(client, session, team.id, kb_id, "制度规定：申请须提前三日。")
    analysis = _analysis_report(client, team.id, kb_id, document["id"])
    assert _run_until_finished(session, analysis["id"]) == "completed"
    base = f"/api/teams/{team.id}/knowledge-bases/{kb_id}/reports"
    report = client.post(f"{base}/from-analysis/{analysis['id']}", json={}).json()
    assert client.delete(f"/api/teams/{team.id}/documents/{document['id']}").status_code == 204

    reopened = client.get(f"{base}/{report['id']}")
    assert reopened.status_code == 200
    body = reopened.json()
    assert body["sections"]
    assert body["sections"][0]["body"]
    source = next(item for item in body["sources"] if item["source_type"] == "document")
    assert source["available"] is False


def test_report_rejects_invalid_analysis_and_other_users_cannot_read(client, session):
    _, team, kb_id = _team_with_kb(client, session, "report-owner@example.com")
    document = _upload_ready(client, session, team.id, kb_id, "制度正文。")
    analysis = _analysis_report(client, team.id, kb_id, document["id"])
    base = f"/api/teams/{team.id}/knowledge-bases/{kb_id}/reports"
    still_running = client.post(f"{base}/from-analysis/{analysis['id']}", json={})
    assert still_running.status_code == 409

    assert _run_until_finished(session, analysis["id"]) == "completed"
    report = client.post(f"{base}/from-analysis/{analysis['id']}", json={}).json()
    other = make_user(session, "report-other@example.com")
    add_member(session, team, other, role="admin")
    client.post("/api/auth/login", json={"email": other.email, "password": "password123"})
    assert client.get(f"{base}/{report['id']}").status_code == 404
    assert client.get(base).json()["items"] == []
    assert client.post(f"{base}/from-analysis/{analysis['id']}", json={}).status_code == 404


def test_report_from_documents_and_notes_plus_reading_helpers(client, session):
    _, team, kb_id = _team_with_kb(client, session)
    document = _upload_ready(
        client,
        session,
        team.id,
        kb_id,
        "适用对象：新员工。\n\n报到时须携带身份证明。",
    )
    note = client.post(
        f"/api/teams/{team.id}/knowledge-bases/{kb_id}/notes",
        json={
            "title": "报到提醒",
            "content": "需要提醒新员工带身份证明。",
            "source_type": "document",
            "source_id": document["id"],
            "citations": [
                {
                    "document_id": document["id"],
                    "document_title": "policy",
                    "locator_label": "第 2 段",
                    "excerpt": "报到时须携带身份证明。",
                }
            ],
        },
    )
    assert note.status_code == 201
    report = client.post(
        f"/api/teams/{team.id}/knowledge-bases/{kb_id}/reports/from-sources",
        json={
            "title": "入职材料说明",
            "purpose": "帮助新员工准备材料",
            "focus": "适用对象与携带材料",
            "length": "short",
            "document_ids": [document["id"]],
            "note_ids": [note.json()["id"]],
        },
    )
    assert report.status_code == 201, report.text
    assert {item["source_type"] for item in report.json()["sources"]} == {
        "document",
        "note",
    }

    guide = client.get(
        f"/api/teams/{team.id}/knowledge-bases/{kb_id}/documents/{document['id']}/guide"
    )
    assert guide.status_code == 200
    assert guide.json()["summary"]
    assert guide.json()["suggested_questions"]
    assert "覆盖已解析" in guide.json()["coverage_note"]

    search = client.get(
        f"/api/teams/{team.id}/knowledge-bases/{kb_id}/document-search",
        params={"q": "身份证明"},
    )
    assert search.status_code == 200
    assert search.json()["hits"][0]["document_id"] == document["id"]
    assert search.json()["hits"][0]["locator_label"]


def test_analysis_suggestion_and_saved_requirements(client, session):
    _, team, kb_id = _team_with_kb(client, session)
    document = _upload_ready(client, session, team.id, kb_id, "合同有效期为 2027 年 12 月。")
    response = client.post(
        f"/api/teams/{team.id}/knowledge-bases/{kb_id}/analysis-suggestions",
        json={"requirement": "比较合同的生效日期和限制条件", "document_ids": [document["id"]]},
    )
    assert response.status_code == 200, response.text
    assert response.json()["fields"]
    assert any(
        field["value_type"] == "date" or "时间" in field["name"]
        for field in response.json()["fields"]
    )

    requirement = client.post(
        f"/api/teams/{team.id}/knowledge-bases/{kb_id}/requirements",
        json={
            "kind": "analysis_fields",
            "name": "合同核对",
            "payload": {"fields": response.json()["fields"]},
        },
    )
    assert requirement.status_code == 201
    listed = client.get(
        f"/api/teams/{team.id}/knowledge-bases/{kb_id}/requirements",
        params={"kind": "analysis_fields"},
    )
    assert listed.status_code == 200
    assert listed.json()["items"][0]["name"] == "合同核对"


def test_cross_knowledge_base_report_sources_are_rejected(client, session):
    _, team, kb_id = _team_with_kb(client, session)
    other_kb_id = client.post(
        f"/api/teams/{team.id}/knowledge-bases", json={"name": "另一库"}
    ).json()["id"]
    document = _upload_ready(client, session, team.id, kb_id, "当前库正文。")
    response = client.post(
        f"/api/teams/{team.id}/knowledge-bases/{other_kb_id}/reports/from-sources",
        json={
            "title": "越界报告",
            "document_ids": [document["id"]],
            "note_ids": [],
        },
    )
    assert response.status_code == 422


def test_pasted_text_and_web_import_enter_processing_pipeline(client, session, monkeypatch):
    _, team, kb_id = _team_with_kb(client, session)
    base = f"/api/teams/{team.id}/documents"
    pasted = client.post(
        f"{base}/text",
        json={
            "knowledge_base_id": kb_id,
            "title": "粘贴的制度摘录",
            "content": "申请差旅须提前三日。",
        },
    )
    assert pasted.status_code == 201, pasted.text
    assert pasted.json()["filename"].endswith(".txt")
    assert pasted.json()["status"] == "pending"

    @contextmanager
    def fake_stream(*_args, **_kwargs):
        yield httpx.Response(
            200,
            headers={"content-type": "text/html; charset=utf-8"},
            content=(
                b"<html><head><title>Public guide</title></head>"
                b"<body><p>Imported public body.</p></body></html>"
            ),
            request=httpx.Request("GET", "https://example.com/guide"),
        )

    monkeypatch.setattr("app.ingestion.httpx.stream", fake_stream)
    imported = client.post(
        f"{base}/web",
        json={
            "knowledge_base_id": kb_id,
            "url": "https://example.com/guide",
        },
    )
    assert imported.status_code == 201, imported.text
    assert imported.json()["title"] == "Public guide"
    assert imported.json()["source_url"] == "https://example.com/guide"
    assert imported.json()["status"] == "pending"

    from app.models import Document, DocumentChunk
    from app.worker import run_once

    for _ in range(5):
        run_once()
        session.expire_all()
        document = session.get(Document, imported.json()["id"])
        if document is not None and document.status == "ready":
            break
    else:
        raise AssertionError(
            f"网页资料未就绪：{document.status if document else 'missing'}"
            f"；{document.failure_reason if document else ''}"
        )
    assert document.failure_reason is None
    assert any(
        "Imported public body." in chunk.content
        for chunk in session.scalars(
            select(DocumentChunk).where(DocumentChunk.document_id == document.id)
        )
    )


def test_replace_document_reports_affected_personal_products(client, session):
    _, team, kb_id = _team_with_kb(client, session)
    document = _upload_ready(
        client,
        session,
        team.id,
        kb_id,
        "旧制度：申请须提前三日。",
        filename="policy.txt",
    )
    note = client.post(
        f"/api/teams/{team.id}/knowledge-bases/{kb_id}/notes",
        json={
            "title": "旧制度笔记",
            "content": "当前申请期限是三日。",
            "source_type": "document",
            "source_id": document["id"],
            "citations": [],
        },
    )
    assert note.status_code == 201
    analysis = _analysis_report(client, team.id, kb_id, document["id"])
    assert _run_until_finished(session, analysis["id"]) == "completed"
    report = client.post(
        f"/api/teams/{team.id}/knowledge-bases/{kb_id}/reports/from-analysis/{analysis['id']}",
        json={},
    )
    assert report.status_code == 201

    replaced = client.post(
        f"/api/teams/{team.id}/documents/{document['id']}/replace",
        data={"knowledge_base_id": kb_id},
        files={"file": ("policy.txt", "新制度：申请须提前七日。", "text/plain")},
    )
    assert replaced.status_code == 201, replaced.text
    body = replaced.json()
    assert body["document"]["id"] != document["id"]
    assert body["document"]["replaces_id"] == document["id"]
    assert body["document"]["title"] == document["title"]
    assert body["affected"]["note_count"] == 1
    assert body["affected"]["analysis_count"] == 1
    assert body["affected"]["report_count"] == 1
    report_detail = client.get(
        f"/api/teams/{team.id}/knowledge-bases/{kb_id}/reports/{report.json()['id']}"
    )
    assert report_detail.status_code == 200
    detail = report_detail.json()
    assert detail["source_review_required"] is True
    assert "已有新版本" in detail["source_review_note"]
    untouched = client.put(
        f"/api/teams/{team.id}/knowledge-bases/{kb_id}/reports/{report.json()['id']}",
        json={
            "title": detail["title"],
            "purpose": detail["purpose"],
            "focus": detail["focus"],
            "length": detail["length"],
            "status": detail["status"],
            "source_review_acknowledged": False,
            "gaps": detail["gaps"],
            "sections": [
                {
                    "id": section["id"],
                    "heading": section["heading"],
                    "body": section["body"],
                }
                for section in detail["sections"]
            ],
        },
    )
    assert untouched.status_code == 200
    assert untouched.json()["source_review_required"] is True
    confirmed = client.put(
        f"/api/teams/{team.id}/knowledge-bases/{kb_id}/reports/{report.json()['id']}",
        json={
            "title": detail["title"],
            "purpose": detail["purpose"],
            "focus": detail["focus"],
            "length": detail["length"],
            "status": detail["status"],
            "source_review_acknowledged": True,
            "gaps": detail["gaps"],
            "sections": [
                {
                    "id": section["id"],
                    "heading": section["heading"],
                    "body": section["body"],
                }
                for section in detail["sections"]
            ],
        },
    )
    assert confirmed.status_code == 200
    assert confirmed.json()["source_review_required"] is False
    assert confirmed.json()["source_review_note"] is None
    assert confirmed.json()["sections"][0]["body"] == detail["sections"][0]["body"]


def test_report_gaps_are_saved_and_section_suggestion_is_preview_only(client, session):
    _, team, kb_id = _team_with_kb(client, session)
    document = _upload_ready(
        client,
        session,
        team.id,
        kb_id,
        "差旅制度：住宿上限为每晚 500 元。",
    )
    analysis = _analysis_report(client, team.id, kb_id, document["id"])
    assert _run_until_finished(session, analysis["id"]) == "completed"
    base = f"/api/teams/{team.id}/knowledge-bases/{kb_id}/reports"
    report = client.post(
        f"{base}/from-analysis/{analysis['id']}",
        json={},
    ).json()
    gap = {"kind": "review", "text": "已人工核对住宿标准。"}
    saved = client.put(
        f"{base}/{report['id']}",
        json={
            "title": report["title"],
            "purpose": report["purpose"],
            "focus": report["focus"],
            "length": report["length"],
            "status": "draft",
            "gaps": [gap],
            "sections": [
                {
                    "id": section["id"],
                    "heading": section["heading"],
                    "body": section["body"],
                }
                for section in report["sections"]
            ],
        },
    )
    assert saved.status_code == 200, saved.text
    assert saved.json()["gaps"] == [gap]
    section = saved.json()["sections"][0]
    original = section["body"]
    suggestion = client.post(
        f"{base}/{report['id']}/sections/{section['id']}/suggest",
        json={"instruction": "保留事实并补上核对提示"},
    )
    assert suggestion.status_code == 200, suggestion.text
    assert suggestion.json()["suggestion"]
    reopened = client.get(f"{base}/{report['id']}").json()
    assert reopened["sections"][0]["body"] == original


def test_source_report_uses_real_model_and_rejects_unbound_citations(
    client, session, monkeypatch
):
    from app.config import get_settings

    _, team, kb_id = _team_with_kb(client, session)
    document = _upload_ready(
        client,
        session,
        team.id,
        kb_id,
        "适用范围：全体员工。\n\n报销期限：费用发生后三十日内提交。",
    )
    settings = get_settings()
    monkeypatch.setattr(settings, "analysis_model_base_url", "https://model.example/v1")
    monkeypatch.setattr(settings, "analysis_model_api_key", "test-key")

    def fake_post(*_args, **kwargs):
        prompt = kwargs["json"]["messages"][1]["content"]
        assert "适用范围：全体员工" in prompt
        return httpx.Response(
            200,
            json={
                "choices": [
                    {
                        "message": {
                            "content": (
                                '{"sections":[{"heading":"适用范围","body":"适用于全体员工。",'
                                '"citation_numbers":[1]},{"heading":"提交期限",'
                                '"body":"报销应在三十日内提交。","citation_numbers":[999]}]}'
                            )
                        }
                    }
                ]
            },
            request=httpx.Request("POST", "https://model.example/v1/chat/completions"),
        )

    monkeypatch.setattr("app.knowledge_products.httpx.post", fake_post)
    response = client.post(
        f"/api/teams/{team.id}/knowledge-bases/{kb_id}/reports/from-sources",
        json={
            "title": "报销说明",
            "purpose": "供员工阅读",
            "focus": "范围与期限",
            "length": "short",
            "document_ids": [document["id"]],
            "note_ids": [],
        },
    )
    assert response.status_code == 201, response.text
    report = response.json()
    assert [section["heading"] for section in report["sections"]] == [
        "适用范围",
        "提交期限",
    ]
    assert len(report["sections"][0]["citations"]) == 1
    assert report["sections"][1]["citations"] == []
