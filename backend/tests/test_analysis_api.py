"""T02：单文档分析的 API + worker 闭环。

这些测试使用真实独立 PostgreSQL、受限运行角色与确定性替身，验证创建 → 领取 →
发布 → 读取、幂等、权限隔离、来源失效与租约竞争。它们不证明提取质量。
"""

from __future__ import annotations

import io
import uuid

from docx import Document as DocxDocument
from sqlalchemy import select, text

from tests.conftest import add_member, make_team, make_user


def _team_with_kb(client, session, email: str = "owner@example.com", role: str = "admin"):
    """建好团队与知识库；知识库只能管理员创建，因此先以管理员身份创建再降级。"""
    from app.models import TeamMembership

    user = make_user(session, email)
    team = make_team(session, "分析团队")
    add_member(session, team, user, role="admin")
    client.post("/api/auth/login", json={"email": email, "password": "password123"})
    kb = client.post(f"/api/teams/{team.id}/knowledge-bases", json={"name": "分析库"}).json()
    if role != "admin":
        membership = session.scalars(
            select(TeamMembership).where(
                TeamMembership.team_id == team.id, TeamMembership.user_id == user.id
            )
        ).one()
        membership.role = role
        session.commit()
    return user, team, kb["id"]


def _upload_ready(client, session, team_id, kb_id, content: str, filename: str = "policy.txt"):
    from app.worker import run_once

    document = client.post(
        f"/api/teams/{team_id}/documents",
        data={"knowledge_base_id": kb_id},
        files={"file": (filename, content, "text/plain")},
    ).json()
    # run_once 公平轮询解析与分析；循环直到文档就绪。
    for _ in range(5):
        run_once()
        session.expire_all()
        listed = client.get(
            f"/api/teams/{team_id}/documents", params={"knowledge_base_id": kb_id}
        ).json()
        current = next(item for item in listed if item["id"] == document["id"])
        if current["status"] == "ready":
            break
    else:
        raise AssertionError(f"文档未就绪：{current['status']}")
    return document


def _create(client, team_id, kb_id, document_id, *, fields=None, kind="extraction", key=None):
    payload = {
        "kind": kind,
        "document_ids": [document_id],
        "fields": fields
        or [
            {"name": "报销标准", "description": "差旅交通费用报销标准", "value_type": "text"},
            {"name": "住宿上限", "description": "住宿费用上限", "value_type": "number"},
        ],
        "title": "差旅制度提取",
    }
    return client.post(
        f"/api/teams/{team_id}/knowledge-bases/{kb_id}/analyses",
        json=payload,
        headers={"Idempotency-Key": key or str(uuid.uuid4())},
    )


def _run_until_finished(session, task_id, rounds: int = 8):
    from app.worker import run_once

    for _ in range(rounds):
        session.expire_all()
        status = session.execute(
            text("select status from analysis_tasks where id = :task_id"),
            {"task_id": task_id},
        ).scalar_one()
        if status in ("completed", "partial", "failed", "invalidated"):
            return status
        run_once()
    raise AssertionError(f"任务未收敛，当前状态 {status}")


def test_create_requires_ready_single_document(client, session):
    _, team, kb_id = _team_with_kb(client, session)
    document = _upload_ready(
        client, session, team.id, kb_id, "差旅交通费用凭有效票据报销，标准为 60%。"
    )

    response = _create(client, team.id, kb_id, document["id"])
    assert response.status_code == 202, response.text
    body = response.json()
    assert body["status"] == "queued"
    assert body["stage"] == "queued"
    assert body["document_count"] == 1
    assert body["kind"] == "extraction"
    assert body["title"] == "差旅制度提取"


def test_create_accepts_batch_and_rejects_unavailable_sources(client, session):
    _, team, kb_id = _team_with_kb(client, session)
    document = _upload_ready(client, session, team.id, kb_id, "差旅交通费用凭有效票据报销。")

    batch = _create(
        client,
        team.id,
        kb_id,
        document["id"],
    ).json()
    assert batch["document_count"] == 1

    # T05 起两份文档可形成批量提取。
    other = _upload_ready(
        client, session, team.id, kb_id, "第二条制度：住宿上限每晚 500 元。", "b.txt"
    )
    response = client.post(
        f"/api/teams/{team.id}/knowledge-bases/{kb_id}/analyses",
        json={
            "kind": "extraction",
            "document_ids": [document["id"], other["id"]],
            "fields": [{"name": "标准", "description": "说明", "value_type": "text"}],
        },
        headers={"Idempotency-Key": str(uuid.uuid4())},
    )
    assert response.status_code == 202
    assert response.json()["document_count"] == 2

    response = client.post(
        f"/api/teams/{team.id}/knowledge-bases/{kb_id}/analyses",
        json={
            "kind": "comparison",
            "document_ids": [document["id"], other["id"]],
            "fields": [{"name": "标准", "description": "说明", "value_type": "text"}],
        },
        headers={"Idempotency-Key": str(uuid.uuid4())},
    )
    assert response.status_code == 202
    assert response.json()["kind"] == "comparison"

    # 先结束前两项，避免用户活跃任务上限先于来源错误返回。
    session.execute(text("update analysis_tasks set status='completed', stage='finished'"))
    session.commit()
    unknown = client.post(
        f"/api/teams/{team.id}/knowledge-bases/{kb_id}/analyses",
        json={
            "kind": "extraction",
            "document_ids": [str(uuid.uuid4())],
            "fields": [{"name": "标准", "description": "说明", "value_type": "text"}],
        },
        headers={"Idempotency-Key": str(uuid.uuid4())},
    )
    assert unknown.status_code == 422
    assert unknown.json()["detail"]["code"] == "ANALYSIS_SOURCE_UNAVAILABLE"


def test_full_loop_publishes_results_and_citations(client, session):
    _, team, kb_id = _team_with_kb(client, session)
    document = _upload_ready(
        client,
        session,
        team.id,
        kb_id,
        "差旅制度：交通费用凭有效票据报销。\n\n住宿费用上限为每晚 500 元。",
    )
    created = _create(client, team.id, kb_id, document["id"]).json()

    status = _run_until_finished(session, created["id"])
    assert status == "completed"

    detail = client.get(f"/api/teams/{team.id}/knowledge-bases/{kb_id}/analyses/{created['id']}")
    assert detail.status_code == 200
    body = detail.json()
    assert body["task"]["status"] == "completed"
    assert body["task"]["stage"] == "finished"
    assert body["task"]["completed_document_count"] == 1
    assert body["source_valid"] is True
    assert len(body["fields"]) == 2
    assert len(body["documents"]) == 1
    assert body["documents"][0]["status"] == "succeeded"
    assert body["documents"][0]["processed_chunk_count"] >= 1

    assert len(body["results"]) == 1
    result = body["results"][0]
    assert len(result["fields"]) == 2
    assert result["document_id"] == document["id"]
    for field in result["fields"]:
        assert field["status"] in ("extracted", "needs_review", "not_mentioned")
        if field["status"] != "not_mentioned":
            assert field["citations"], "已提取字段必须带依据"
            citation = field["citations"][0]
            assert citation["document_id"] == document["id"]
            assert citation["excerpt"]
            assert citation["locator_label"]


def test_full_loop_preserves_docx_table_citation_structure(client, session):
    _, team, kb_id = _team_with_kb(client, session)
    source = DocxDocument()
    source.add_paragraph("差旅标准如下。")
    table = source.add_table(rows=3, cols=3)
    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "标准"
    table.cell(0, 2).text = "备注"
    table.cell(1, 0).text = "住宿费"
    table.cell(1, 1).text = "每晚 500 元"
    table.cell(2, 0).text = "交通费"
    table.cell(2, 1).text = "凭有效票据"
    buffer = io.BytesIO()
    source.save(buffer)
    document = client.post(
        f"/api/teams/{team.id}/documents",
        data={"knowledge_base_id": kb_id},
        files={
            "file": (
                "差旅标准.docx",
                buffer.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    ).json()
    from app.worker import run_once

    for _ in range(8):
        run_once()
        session.expire_all()
        listed = client.get(
            f"/api/teams/{team.id}/documents", params={"knowledge_base_id": kb_id}
        ).json()
        current = next(item for item in listed if item["id"] == document["id"])
        if current["status"] == "ready":
            break
    assert current["status"] == "ready"

    created = _create(
        client,
        team.id,
        kb_id,
        document["id"],
        fields=[
            {
        "name": "住宿费",
        "description": "住宿费",
                "value_type": "text",
            }
        ],
    ).json()
    assert _run_until_finished(session, created["id"]) == "completed"
    detail = client.get(
        f"/api/teams/{team.id}/knowledge-bases/{kb_id}/analyses/{created['id']}"
    ).json()
    citation = detail["results"][0]["fields"][0]["citations"][0]
    assert citation["table_number"] == 1
    assert citation["table_row"] == 2
    assert citation["table_headers"] == ["项目", "标准", "备注"]
    assert citation["table_cells"] == ["住宿费", "每晚 500 元", ""]


def test_idempotent_replay_returns_same_task(client, session):
    _, team, kb_id = _team_with_kb(client, session)
    document = _upload_ready(client, session, team.id, kb_id, "交通费用凭有效票据报销。")
    key = str(uuid.uuid4())

    first = _create(client, team.id, kb_id, document["id"], key=key)
    second = _create(client, team.id, kb_id, document["id"], key=key)
    assert first.status_code == 202
    assert second.status_code == 200
    assert first.json()["id"] == second.json()["id"]

    conflicting = _create(
        client,
        team.id,
        kb_id,
        document["id"],
        key=key,
        fields=[{"name": "另一个字段", "description": "不同说明", "value_type": "text"}],
    )
    assert conflicting.status_code == 409
    assert conflicting.json()["detail"]["code"] == "ANALYSIS_IDEMPOTENCY_CONFLICT"


def test_other_member_and_admin_cannot_read_others_analysis(client, session):
    user, team, kb_id = _team_with_kb(client, session, "owner@example.com", role="member")
    document = _upload_ready(client, session, team.id, kb_id, "交通费用凭有效票据报销。")
    created = _create(client, team.id, kb_id, document["id"]).json()

    # 同团队另一位管理员：团队资料共享，但分析结果仅创建者可见。
    other = make_user(session, "other@example.com")
    add_member(session, team, other, role="admin")
    client.post("/api/auth/login", json={"email": other.email, "password": "password123"})

    detail = client.get(f"/api/teams/{team.id}/knowledge-bases/{kb_id}/analyses/{created['id']}")
    assert detail.status_code == 404

    listing = client.get(f"/api/teams/{team.id}/knowledge-bases/{kb_id}/analyses")
    assert listing.status_code == 200
    assert listing.json()["items"] == []
    assert (
        client.delete(
            f"/api/teams/{team.id}/knowledge-bases/{kb_id}/analyses/{created['id']}"
        ).status_code
        == 404
    )


def test_delete_analysis_keeps_report_and_citation_snapshot(client, session):
    _, team, kb_id = _team_with_kb(client, session)
    document = _upload_ready(
        client, session, team.id, kb_id, "差旅制度：交通费用凭有效票据报销。"
    )
    analysis = _create(client, team.id, kb_id, document["id"]).json()
    assert _run_until_finished(session, analysis["id"]) == "completed"
    base = f"/api/teams/{team.id}/knowledge-bases/{kb_id}"
    report = client.post(
        f"{base}/reports/from-analysis/{analysis['id']}",
        json={"purpose": "提交审阅", "focus": "交通", "length": "medium"},
    ).json()

    response = client.delete(f"{base}/analyses/{analysis['id']}")
    assert response.status_code == 204
    assert client.get(f"{base}/analyses/{analysis['id']}").status_code == 404
    reopened = client.get(f"{base}/reports/{report['id']}")
    assert reopened.status_code == 200
    body = reopened.json()
    assert body["source_analysis_id"] is None
    assert any(section["citations"] for section in body["sections"])
    assert any(
        source["source_type"] == "analysis"
        and "来源分析已删除" in source["title_snapshot"]
        and source["available"] is False
        for source in body["sources"]
    )


def test_delete_running_analysis_conflicts(client, session):
    _, team, kb_id = _team_with_kb(client, session)
    document = _upload_ready(client, session, team.id, kb_id, "交通费用凭有效票据报销。")
    analysis = _create(client, team.id, kb_id, document["id"]).json()
    response = client.delete(
        f"/api/teams/{team.id}/knowledge-bases/{kb_id}/analyses/{analysis['id']}"
    )
    assert response.status_code == 409
    assert response.json()["detail"]["code"] == "ANALYSIS_ACTIVE"


def test_cross_team_access_is_rejected(client, session):
    _, team_a, kb_a = _team_with_kb(client, session, "a@example.com")
    document = _upload_ready(client, session, team_a.id, kb_a, "交通费用凭有效票据报销。")
    created = _create(client, team_a.id, kb_a, document["id"]).json()

    # 切换到 B 团队用户后再创建 B 的知识库；此时 A 的登录态已被替换。
    other, team_b, kb_b = _team_with_kb(client, session, "b@example.com")

    assert (
        client.get(
            f"/api/teams/{team_b.id}/knowledge-bases/{kb_b}/analyses/{created['id']}"
        ).status_code
        == 404
    )
    # 试图用 B 团队路径访问 A 的知识库也应被拒绝。
    assert (
        client.get(
            f"/api/teams/{team_b.id}/knowledge-bases/{kb_a}/analyses/{created['id']}"
        ).status_code
        == 404
    )


def test_source_deletion_invalidates_and_hides_results(client, session):
    _, team, kb_id = _team_with_kb(client, session)
    document = _upload_ready(
        client, session, team.id, kb_id, "差旅制度：交通费用凭有效票据报销，标准 60%。"
    )
    created = _create(client, team.id, kb_id, document["id"]).json()
    assert _run_until_finished(session, created["id"]) == "completed"

    detail = client.get(
        f"/api/teams/{team.id}/knowledge-bases/{kb_id}/analyses/{created['id']}"
    ).json()
    assert detail["results"]

    assert client.delete(f"/api/teams/{team.id}/documents/{document['id']}").status_code == 204

    after = client.get(f"/api/teams/{team.id}/knowledge-bases/{kb_id}/analyses/{created['id']}")
    assert after.status_code == 200
    body = after.json()
    assert body["task"]["status"] == "invalidated"
    assert body["source_valid"] is False
    assert body["invalid_reason"] == "所选文档已被删除"
    # 字段定义与文档范围保留，但所有取值、摘录与结论消失。
    assert body["fields"]
    assert body["results"] == []
    assert body["conclusions"] == []
    assert "excerpt" not in after.text


def test_source_text_change_invalidates_analysis(client, session):
    _, team, kb_id = _team_with_kb(client, session)
    document = _upload_ready(client, session, team.id, kb_id, "交通费用凭有效票据报销。")
    created = _create(client, team.id, kb_id, document["id"]).json()
    assert _run_until_finished(session, created["id"]) == "completed"

    # 模拟重索引后正文变化：片段被替换。
    session.execute(
        text("update document_chunks set content = :content where document_id = :doc_id"),
        {"content": "已改写的制度正文。", "doc_id": document["id"]},
    )
    session.commit()

    body = client.get(
        f"/api/teams/{team.id}/knowledge-bases/{kb_id}/analyses/{created['id']}"
    ).json()
    assert body["task"]["status"] == "invalidated"
    assert body["source_valid"] is False
    assert body["results"] == []


def test_member_removal_stops_publication(client, session):
    user, team, kb_id = _team_with_kb(client, session, "member@example.com", role="member")
    document = _upload_ready(client, session, team.id, kb_id, "交通费用凭有效票据报销。")
    created = _create(client, team.id, kb_id, document["id"]).json()

    from app.models import TeamMembership

    membership = session.scalars(
        select(TeamMembership).where(
            TeamMembership.team_id == team.id, TeamMembership.user_id == user.id
        )
    ).one()
    session.delete(membership)
    session.commit()

    status = _run_until_finished(session, created["id"])
    assert status == "failed"
    session.expire_all()
    row = session.execute(
        text("select error_code from analysis_tasks where id = :id"), {"id": created["id"]}
    ).one()
    assert row.error_code == "ANALYSIS_PERMISSION_REVOKED"

    from app.models import AnalysisFieldResult

    assert (
        session.scalars(
            select(AnalysisFieldResult).where(
                AnalysisFieldResult.task_id == uuid.UUID(created["id"])
            )
        ).all()
        == []
    )


def test_stale_lease_cannot_publish(client, session):
    from app.analysis_tasks import AnalysisLeaseLost
    from app.analysis_worker import run_analysis_document
    from app.db import apply_context
    from app.models import AnalysisDocument

    _, team, kb_id = _team_with_kb(client, session)
    document = _upload_ready(client, session, team.id, kb_id, "交通费用凭有效票据报销。")
    _create(client, team.id, kb_id, document["id"])

    from app.analysis_tasks import claim_next_document

    claimed = claim_next_document(session, 60)
    assert claimed is not None
    session.commit()

    # 另一个执行者接管：清空旧令牌并重新领取，得到新令牌。
    session.execute(
        text(
            "update analysis_documents set status='queued', lease_token=null, "
            "lease_expires_at=null where id = :id"
        ),
        {"id": claimed.document_execution_id},
    )
    session.commit()
    second = claim_next_document(session, 60)
    assert second is not None
    session.commit()
    assert second.lease_token != claimed.lease_token

    apply_context(session, team_id=team.id, user_id=claimed.owner_id)
    with __import__("pytest").raises(AnalysisLeaseLost):
        run_analysis_document(
            session,
            execution_id=claimed.document_execution_id,
            task_id=claimed.task_id,
            team_id=claimed.team_id,
            owner_id=claimed.owner_id,
            document_id=claimed.document_id,
            lease_token=claimed.lease_token,
        )
    session.rollback()

    session.expire_all()
    execution = session.get(AnalysisDocument, claimed.document_execution_id)
    assert execution is not None and execution.status == "running"
    from app.models import AnalysisFieldResult

    assert (
        session.scalars(
            select(AnalysisFieldResult).where(
                AnalysisFieldResult.document_execution_id == claimed.document_execution_id
            )
        ).all()
        == []
    )


def test_list_is_paginated_and_owner_only(client, session):
    _, team, kb_id = _team_with_kb(client, session)
    document = _upload_ready(client, session, team.id, kb_id, "交通费用凭有效票据报销。")

    for index in range(3):
        response = _create(
            client,
            team.id,
            kb_id,
            document["id"],
            fields=[
                {
                    "name": f"字段{index}",
                    "description": "说明",
                    "value_type": "text",
                }
            ],
        )
        assert response.status_code == 202, response.text
        # 每个用户最多两个活跃任务；分页测试逐个收敛，避免撞上并发上限。
        session.execute(
            text(
                "update analysis_tasks set status='completed', stage='finished' where id = :task_id"
            ),
            {"task_id": response.json()["id"]},
        )
        session.commit()

    first = client.get(
        f"/api/teams/{team.id}/knowledge-bases/{kb_id}/analyses", params={"limit": 2}
    ).json()
    assert len(first["items"]) == 2
    assert first["next_cursor"]

    second = client.get(
        f"/api/teams/{team.id}/knowledge-bases/{kb_id}/analyses",
        params={"limit": 2, "cursor": first["next_cursor"]},
    ).json()
    assert len(second["items"]) == 1
    assert second["next_cursor"] is None

    ids = {item["id"] for item in first["items"] + second["items"]}
    assert len(ids) == 3
    # 列表只含元数据，不泄漏字段值或摘录。
    assert "results" not in first["items"][0]
    assert "fields" not in first["items"][0]


def test_bad_cursor_and_missing_idempotency_key(client, session):
    _, team, kb_id = _team_with_kb(client, session)
    document = _upload_ready(client, session, team.id, kb_id, "交通费用凭有效票据报销。")

    bad = client.get(
        f"/api/teams/{team.id}/knowledge-bases/{kb_id}/analyses", params={"cursor": "!!!"}
    )
    assert bad.status_code == 422
    assert bad.json()["detail"]["code"] == "ANALYSIS_INVALID_INPUT"

    missing = client.post(
        f"/api/teams/{team.id}/knowledge-bases/{kb_id}/analyses",
        json={
            "kind": "extraction",
            "document_ids": [document["id"]],
            "fields": [{"name": "标准", "description": "说明", "value_type": "text"}],
        },
        headers={"Idempotency-Key": ""},
    )
    assert missing.status_code == 422


def test_rls_context_does_not_leak_across_pooled_connections(client, session):
    """连接池复用后不能残留上一请求的团队/用户上下文。"""
    from app.db import apply_context
    from tests.conftest import AppSession

    _, team, kb_id = _team_with_kb(client, session)
    document = _upload_ready(client, session, team.id, kb_id, "交通费用凭有效票据报销。")
    created = _create(client, team.id, kb_id, document["id"]).json()
    assert _run_until_finished(session, created["id"]) == "completed"

    worker = AppSession()
    try:
        apply_context(worker, team_id=team.id, user_id=None)
        visible = worker.execute(text("select count(*) from analysis_tasks")).scalar_one()
        # 只有团队上下文、没有用户上下文时 RLS 必须拒绝返回私人行。
        assert visible == 0
        worker.commit()

        apply_context(worker, team_id=team.id, user_id=uuid.uuid4())
        assert worker.execute(text("select count(*) from analysis_tasks")).scalar_one() == 0
        worker.commit()
    finally:
        worker.close()

    # 同一个应用请求仍能读到自己的任务，说明上下文没有在池中被污染。
    body = client.get(f"/api/teams/{team.id}/knowledge-bases/{kb_id}/analyses/{created['id']}")
    assert body.status_code == 200


def test_analysis_task_survives_worker_restart(client, session):
    """worker 进程重启后，未完成的任务仍可被重新领取并完成。"""
    _, team, kb_id = _team_with_kb(client, session)
    document = _upload_ready(client, session, team.id, kb_id, "交通费用凭有效票据报销。")
    created = _create(client, team.id, kb_id, document["id"]).json()

    # 模拟 worker 在领取后崩溃：执行停在 running，租约过期。
    from datetime import UTC, datetime, timedelta

    from app.analysis_tasks import claim_next_document

    claimed = claim_next_document(session, 60)
    assert claimed is not None
    session.commit()
    session.execute(
        text("update analysis_documents set lease_expires_at = :past where id = :id"),
        {
            "past": datetime.now(UTC) - timedelta(seconds=5),
            "id": claimed.document_execution_id,
        },
    )
    session.commit()

    status = _run_until_finished(session, created["id"])
    assert status == "completed"
    session.expire_all()
    assert (
        session.execute(
            text("select status from analysis_documents where id = :id"),
            {"id": claimed.document_execution_id},
        ).scalar_one()
        == "succeeded"
    )


def test_upload_parsing_is_not_starved_by_analysis(client, session):
    """两条路径公平轮询：分析排队时上传解析仍能在有限轮次内完成。"""
    from app.worker import run_once

    _, team, kb_id = _team_with_kb(client, session)
    ready = _upload_ready(client, session, team.id, kb_id, "交通费用凭有效票据报销。")
    _create(client, team.id, kb_id, ready["id"])

    pending = client.post(
        f"/api/teams/{team.id}/documents",
        data={"knowledge_base_id": kb_id},
        files={"file": ("second.txt", "第二条候选资料正文。", "text/plain")},
    ).json()
    assert pending["status"] == "pending"

    for _ in range(4):
        run_once()

    session.expire_all()
    listed = client.get(
        f"/api/teams/{team.id}/documents", params={"knowledge_base_id": kb_id}
    ).json()
    current = next(item for item in listed if item["id"] == pending["id"])
    assert current["status"] == "ready"


def test_active_task_limit_is_enforced(client, session):
    _, team, kb_id = _team_with_kb(client, session)
    document = _upload_ready(client, session, team.id, kb_id, "交通费用凭有效票据报销。")

    responses = [
        _create(
            client,
            team.id,
            kb_id,
            document["id"],
            fields=[{"name": f"字段{index}", "description": "说明", "value_type": "text"}],
        )
        for index in range(3)
    ]
    assert [item.status_code for item in responses] == [202, 202, 409]
    assert responses[2].json()["detail"]["code"] == "ANALYSIS_LIMIT_EXCEEDED"
